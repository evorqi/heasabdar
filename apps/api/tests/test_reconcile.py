from __future__ import annotations

import hashlib
from decimal import Decimal
from io import BytesIO

import openpyxl
from fastapi.testclient import TestClient

from tests.test_api import create_workspace, upload

SHEETS = ["خلاصه", "منطبق", "مغایر", "فقط فایل اول", "فقط فایل دوم", "مشکوک", "درباره این خروجی"]


def _xlsx(rows: list[list]) -> bytes:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "فروش"
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    book.save(buffer)
    book.close()
    return buffer.getvalue()


def _pair(client: TestClient) -> tuple[dict, dict, dict, bytes, bytes]:
    left = _xlsx(
        [
            ["شماره فاکتور", "نام", "مبلغ"],
            ["1001", "علی", 1000],
            ["1002", "سارا", 0],
            ["1003", "رضا", 500],
            [None, "نیکان", 200],
        ]
    )
    right = _xlsx(
        [
            ["شماره فاکتور", "نام", "مبلغ"],
            ["۱۰۰۱", "علي", 1000],
            ["1002", "سارا", 50],
            ["1004", "مریم", 300],
            [None, "نيكان", 200],
        ]
    )
    workspace = create_workspace(client, "مغایرت فروش")
    uploaded = upload(
        client,
        workspace["id"],
        [("فروش اول.xlsx", left), ("فروش دوم.xlsx", right)],
    )
    assert uploaded.status_code == 201, uploaded.text
    accepted = {item["original_name"]: item for item in uploaded.json()["accepted"]}
    return workspace, accepted["فروش اول.xlsx"], accepted["فروش دوم.xlsx"], left, right


def test_two_files_reconcile_from_upload_through_reread(client: TestClient) -> None:
    workspace, left, right, left_bytes, right_bytes = _pair(client)
    preview = client.post(
        f"/api/workspaces/{workspace['id']}/reconciliations/preview",
        json={"left_file_id": left["id"], "right_file_id": right["id"]},
    )
    assert preview.status_code == 200, preview.text
    assert preview.json()["key_ambiguous"] is False
    assert preview.json()["key_columns"] == ["شماره فاکتور"]
    assert preview.json()["result_file_id"] is None
    assert preview.json()["status"] == "preview"

    applied = client.post(
        f"/api/workspaces/{workspace['id']}/reconciliations",
        json={"left_file_id": left["id"], "right_file_id": right["id"]},
    )
    assert applied.status_code == 201, applied.text
    report = applied.json()
    assert report["status"] == "succeeded"
    assert report["key_decision"] == "detected"
    assert report["validation"]["ok"] is True
    assert report["validation"]["scale"] == "1"
    rows = report["rows"]
    assert _count(rows, "left") == 4
    assert _count(rows, "right") == 4
    assert _amount(rows, "left") == Decimal("1700")
    assert _amount(rows, "right") == Decimal("1550")
    assert Decimal(report["validation"]["left_amount"]) == Decimal("1700")
    assert Decimal(report["validation"]["right_amount"]) == Decimal("1550")

    matched = _one(rows, "full_match")
    assert matched["key_right"] == "۱۰۰۱"
    assert matched["key_left"] == "1001"
    assert matched["amount_delta_ratio"] == "0"
    assert "۱۰۰۱" in matched["right_values"]["شماره فاکتور"]

    mismatch = _one(rows, "value_mismatch")
    assert mismatch["amount_left"] == "0"
    assert mismatch["amount_right"] == "50"
    assert mismatch["amount_delta"] == "50"
    assert mismatch["amount_delta_ratio"] is None
    assert mismatch["reason"] == "مبلغ مبنای مقایسه صفر است."

    assert _one(rows, "only_left")["key_left"] == "1003"
    assert _one(rows, "only_right")["key_right"] == "1004"
    ambiguous = _one(rows, "ambiguous")
    assert ambiguous["status"] != "full_match"
    assert ambiguous["name_left"] == "نیکان"
    assert ambiguous["name_right"] == "نيكان"
    assert "شناسه مشترک ندارد" in ambiguous["reason"]

    listed = client.get(f"/api/workspaces/{workspace['id']}/files")
    assert listed.status_code == 200
    files = {item["original_name"]: item for item in listed.json()["files"]}
    assert files["فروش اول.xlsx"]["sha256"] == hashlib.sha256(left_bytes).hexdigest()
    assert files["فروش دوم.xlsx"]["sha256"] == hashlib.sha256(right_bytes).hexdigest()
    left_download = client.get(
        f"/api/workspaces/{workspace['id']}/versions/{left['current_version_id']}/download"
    )
    right_download = client.get(
        f"/api/workspaces/{workspace['id']}/versions/{right['current_version_id']}/download"
    )
    assert left_download.content == left_bytes
    assert right_download.content == right_bytes

    output = client.get(
        f"/api/workspaces/{workspace['id']}/versions/{report['result_version_id']}/download"
    )
    assert output.status_code == 200
    book = openpyxl.load_workbook(BytesIO(output.content))
    assert book.sheetnames == SHEETS
    assert all(sheet.sheet_view.rightToLeft for sheet in book.worksheets)
    assert book["خلاصه"]["A2"].value == "منطبق"
    assert book["خلاصه"]["B2"].value == "1"
    assert book["خلاصه"]["B7"].value == "4"
    matched_values = [cell.value for row in book["منطبق"].iter_rows() for cell in row]
    assert "۱۰۰۱" in matched_values
    assert "1001" in matched_values
    suspect_values = [cell.value for row in book["مشکوک"].iter_rows() for cell in row]
    assert "نیکان" in suspect_values
    assert "نيكان" in suspect_values
    book.close()


def test_ambiguous_key_writes_nothing_until_the_user_chooses(client: TestClient) -> None:
    left = _xlsx(
        [
            ["کد ملی", "شماره فاکتور", "مبلغ"],
            ["001", "1001", 10],
            ["002", "1002", 20],
        ]
    )
    right = _xlsx(
        [
            ["کد ملی", "شماره فاکتور", "مبلغ"],
            ["001", "1001", 10],
            ["002", "1002", 20],
        ]
    )
    workspace = create_workspace(client, "کلید مبهم")
    uploaded = upload(client, workspace["id"], [("چپ.xlsx", left), ("راست.xlsx", right)])
    assert uploaded.status_code == 201, uploaded.text
    accepted = {item["original_name"]: item for item in uploaded.json()["accepted"]}
    body = {
        "left_file_id": accepted["چپ.xlsx"]["id"],
        "right_file_id": accepted["راست.xlsx"]["id"],
    }
    preview = client.post(f"/api/workspaces/{workspace['id']}/reconciliations/preview", json=body)
    assert preview.status_code == 200, preview.text
    assert preview.json()["key_ambiguous"] is True
    assert set(preview.json()["key_choices"]) == {"کد ملی", "شماره فاکتور"}
    assert preview.json()["rows"] == []
    assert preview.json()["result_file_id"] is None

    waiting = client.post(f"/api/workspaces/{workspace['id']}/reconciliations", json=body)
    assert waiting.status_code == 200, waiting.text
    assert waiting.json()["status"] == "waiting_confirmation"
    assert waiting.json()["result_file_id"] is None
    assert waiting.json()["key_decision"] is None
    listed = client.get(f"/api/workspaces/{workspace['id']}/files")
    assert {item["original_name"] for item in listed.json()["files"]} == {"چپ.xlsx", "راست.xlsx"}

    chosen = client.post(
        f"/api/workspaces/{workspace['id']}/reconciliations",
        json={**body, "key_columns": ["شماره فاکتور"]},
    )
    assert chosen.status_code == 201, chosen.text
    report = chosen.json()
    assert report["key_decision"] == "confirmed_by_user"
    assert report["key_columns"] == ["شماره فاکتور"]
    assert report["summary"]["full_match"] == 2
    assert report["result_file_id"]
    names = {item["original_name"] for item in client.get(f"/api/workspaces/{workspace['id']}/files").json()["files"]}
    assert "مغایرت.xlsx" in names

    missing = client.post(
        f"/api/workspaces/{workspace['id']}/reconciliations",
        json={**body, "key_columns": ["شهر"]},
    )
    assert missing.status_code == 400
    assert missing.json()["code"] == "request.invalid"
    assert missing.json()["message"] == "ستون انتخاب‌شده در هر دو فایل نیست."


def _count(rows: list[dict], side: str) -> int:
    total = 0
    for row in rows:
        if row["status"] in {"full_match", "value_mismatch", f"only_{side}"}:
            total += 1
        elif row["status"] == "ambiguous" and row[f"{side}_values"]:
            total += 1
    return total


def _amount(rows: list[dict], side: str) -> Decimal:
    total = Decimal(0)
    field = f"amount_{side}"
    for row in rows:
        if side == "left" and row["status"] == "only_right":
            continue
        if side == "right" and row["status"] == "only_left":
            continue
        if row[field] is not None:
            total += Decimal(row[field])
    return total


def _one(rows: list[dict], status: str) -> dict:
    found = [row for row in rows if row["status"] == status]
    assert len(found) == 1, found
    return found[0]
