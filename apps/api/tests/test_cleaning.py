from __future__ import annotations

import hashlib
from io import BytesIO

import openpyxl
from fastapi.testclient import TestClient

from hesabdar.cleaning.rules import clean_cell, national_id_checksum_ok
from hesabdar.config import Settings
from tests.test_api import create_workspace, upload


def test_letters_digits_dates_and_identifiers_follow_their_rules() -> None:
    final, changes, skipped = clean_cell(
        "  علي كريمي  ", header=False, kind="text", role="person_name", formula=False, error_token=None
    )
    assert final == "علی کریمی"
    assert [item[0] for item in changes] == ["text.trim", "text.arabic_yeh", "text.arabic_kaf"]
    assert skipped == []

    untouched, space_changes, _skipped = clean_cell(
        "علی‌رضا  خان", header=False, kind="text", role=None, formula=False, error_token=None
    )
    assert untouched == "علی‌رضا خان"
    assert "\u200c" in untouched
    assert [item[0] for item in space_changes] == ["text.collapse_spaces"]

    identifier, id_changes, id_skipped = clean_cell(
        "۰۰۱۲۳", header=False, kind="identifier", role="invoice_no", formula=False, error_token=None
    )
    assert identifier == "00123"
    assert identifier != "123"
    assert id_changes == [("identifier.leading_zeros", "۰۰۱۲۳", "00123")]
    assert id_skipped == []

    number, number_changes, number_skipped = clean_cell(
        "۱۲۳۴", header=False, kind="number", role="amount", formula=False, error_token=None
    )
    assert number == "1234"
    assert number_changes[-1][0] == "number.parse"
    assert number_changes[-1][1] == "۱۲۳۴"
    assert number_skipped == []

    ambiguous_number, _changes, skips = clean_cell(
        "1,234", header=False, kind="number", role="amount", formula=False, error_token=None
    )
    assert ambiguous_number == "1,234"
    assert skips[0][0] == "number.ambiguous"

    ambiguous_date, _changes, date_skips = clean_cell(
        "02/03/04", header=False, kind="date", role="date", formula=False, error_token=None
    )
    assert ambiguous_date == "02/03/04"
    assert date_skips[0][0] == "date.ambiguous"

    invalid, _changes, invalid_skips = clean_cell(
        "1403/13/01", header=False, kind="date", role="date", formula=False, error_token=None
    )
    assert invalid == "1403/13/01"
    assert invalid_skips[0][0] == "date.invalid"

    prose, _changes, prose_skips = clean_cell(
        "ندارد", header=False, kind="number", role="amount", formula=False, error_token=None
    )
    assert prose == "ندارد"
    assert prose_skips[0][0] == "number.not_numeric"
    assert national_id_checksum_ok("0012345678") is False
    kept, kept_changes, kept_skips = clean_cell(
        "0012345678",
        header=False,
        kind="identifier",
        role="national_id",
        formula=False,
        error_token=None,
    )
    assert kept == "0012345678"
    assert kept_changes == []
    assert kept_skips[0][0] == "identifier.national_id_checksum"


def test_cleaning_writes_a_child_and_keeps_the_original(
    client: TestClient, settings: Settings
) -> None:
    original = _acceptance_workbook()
    workspace = create_workspace(client, "پاک‌سازی")
    stored = _store(client, workspace["id"], "مشتری‌ها.xlsx", original)
    preview = client.post(
        f"/api/workspaces/{workspace['id']}/cleanings/preview",
        json={"file_id": stored["id"]},
    )
    assert preview.status_code == 200, preview.text
    body = preview.json()
    assert body["stats"]["change_count"] == len(body["changes"])
    assert body["result_version_id"] is None
    assert any(item["rule_id"] == "text.arabic_yeh" for item in body["changes"])
    assert any(item["rule_id"] == "text.arabic_kaf" for item in body["changes"])
    assert any(item["after"] == "00123" and item["before"] == "۰۰۱۲۳" for item in body["changes"])
    assert any(item["rule_id"] == "number.parse" and item["before"] == "۱۲۳۴" for item in body["changes"])
    assert any(item["rule_id"] == "date.ambiguous" for item in body["skipped"])
    assert any(item["rule_id"] == "date.invalid" for item in body["skipped"])
    assert any(item["rule_id"] == "number.ambiguous" for item in body["skipped"])
    assert any(item["rule_id"] == "column.empty_required" for item in body["skipped"])
    assert any(item["rule_id"] == "identifier.national_id_checksum" for item in body["skipped"])

    waiting = client.post(
        f"/api/workspaces/{workspace['id']}/cleanings",
        json={"file_id": stored["id"], "confirm": False},
    )
    assert waiting.status_code == 200, waiting.text
    assert waiting.json()["status"] == "waiting_confirmation"
    assert waiting.json()["result_version_id"] is None
    assert _blob(settings, workspace["id"], stored["current_version_id"]) == original

    applied = client.post(
        f"/api/workspaces/{workspace['id']}/cleanings",
        json={"file_id": stored["id"], "confirm": True},
    )
    assert applied.status_code == 201, applied.text
    report = applied.json()
    assert report["status"] == "succeeded"
    assert report["stats"]["change_count"] == len(report["changes"])
    assert report["message"]
    result_id = report["result_version_id"]
    assert result_id
    assert _blob(settings, workspace["id"], stored["current_version_id"]) == original

    cleaned = client.get(f"/api/workspaces/{workspace['id']}/versions/{result_id}/download")
    assert cleaned.status_code == 200
    assert cleaned.content != original
    book = openpyxl.load_workbook(BytesIO(cleaned.content))
    sheet = book["مشتری‌ها"]
    assert sheet.sheet_view.rightToLeft is True
    assert sheet["A2"].value == "علی کریمی"
    assert sheet["B2"].value == "00123"
    assert sheet["C2"].value == "1234"
    assert sheet["D2"].value == "1403-07-21"
    assert sheet["A3"].value == "علی‌رضا"
    assert "\u200c" in sheet["A3"].value
    assert sheet["B3"].value == "0012345678"
    assert sheet["C3"].value == "ندارد"
    assert sheet["D3"].value == "02/03/04"
    assert sheet["C4"].value == "1,234"
    assert sheet["D4"].value == "1403/13/01"
    assert sheet["E2"].value == "=cmd|'/c calc'!A0"
    assert sheet["E2"].data_type != "f"
    assert sheet["A5"].value == "تکرار"
    assert sheet["A6"].value == "تکرار"
    book.close()

    restored = client.post(
        f"/api/workspaces/{workspace['id']}/files/{stored['id']}/restore",
        json={"version_id": stored["current_version_id"]},
    )
    assert restored.status_code == 200, restored.text
    assert restored.json()["current_version_id"] == stored["current_version_id"]
    again = client.get(
        f"/api/workspaces/{workspace['id']}/versions/{stored['current_version_id']}/download"
    )
    assert again.content == original
    still = client.get(f"/api/workspaces/{workspace['id']}/versions/{result_id}/download")
    assert hashlib.sha256(still.content).hexdigest() == hashlib.sha256(cleaned.content).hexdigest()


def test_csv_neutralizes_formula_text_and_keeps_persian(client: TestClient) -> None:
    payload = "شرح,مبلغ\n=cmd|'/c calc'!A0,۱۲۳\n".encode("utf-8-sig")
    workspace = create_workspace(client, "csv")
    stored = _store(client, workspace["id"], "لیست.csv", payload)
    applied = client.post(
        f"/api/workspaces/{workspace['id']}/cleanings",
        json={"file_id": stored["id"], "confirm": True},
    )
    assert applied.status_code == 201, applied.text
    report = applied.json()
    assert report["stats"]["change_count"] == len(report["changes"])
    assert any(item["rule_id"] == "output.neutralize" for item in report["changes"])
    assert any(item["rule_id"] == "number.parse" and item["before"] == "۱۲۳" for item in report["changes"])
    downloaded = client.get(
        f"/api/workspaces/{workspace['id']}/versions/{report['result_version_id']}/download"
    )
    text = downloaded.content.decode("utf-8-sig")
    assert text.startswith("شرح,مبلغ")
    assert "'=cmd|'/c calc'!A0" in text
    assert "123" in text


def test_small_change_does_not_require_confirmation(client: TestClient, settings: Settings) -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "فروش"
    sheet.append(["نام"] + [""] * 0)
    sheet["A1"] = "نام"
    for index in range(1, 6):
        sheet.cell(index + 1, 1, "سارا" if index > 1 else "علي")
    original = _save(book)
    workspace = create_workspace(client)
    stored = _store(client, workspace["id"], "کم.xlsx", original)
    applied = client.post(
        f"/api/workspaces/{workspace['id']}/cleanings",
        json={"file_id": stored["id"]},
    )
    assert applied.status_code == 201, applied.text
    assert applied.json()["status"] == "succeeded"
    assert _blob(settings, workspace["id"], stored["current_version_id"]) == original


def _acceptance_workbook() -> bytes:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "مشتری‌ها"
    sheet.append(["نام", "شماره فاکتور", "مبلغ", "تاریخ", "شرح", "کد ملی"])
    sheet.append(["علي كريمي", "۰۰۱۲۳", "۱۲۳۴", "۱۴۰۳/۰۷/۲۱", "=cmd|'/c calc'!A0", ""])
    sheet["E2"].data_type = "s"
    sheet["E2"].number_format = "@"
    sheet.append(["علی‌رضا", "0012345678", "ندارد", "02/03/04", None, "0012345678"])
    sheet.append(["سارا", "100", "1,234", "1403/13/01", None, "0123456789"])
    sheet.append(["تکرار", "200", "10", "1403/01/01", None, None])
    sheet.append(["تکرار", "200", "10", "1403/01/01", None, None])
    return _save(book)


def _save(book: openpyxl.Workbook) -> bytes:
    buffer = BytesIO()
    book.save(buffer)
    book.close()
    return buffer.getvalue()


def _store(client: TestClient, workspace_id: str, name: str, data: bytes) -> dict:
    response = upload(client, workspace_id, [(name, data)])
    assert response.status_code == 201, response.text
    return response.json()["accepted"][0]


def _blob(settings: Settings, workspace_id: str, version_id: str) -> bytes:
    return (settings.data_dir / "blobs" / workspace_id / version_id).read_bytes()
