from __future__ import annotations

from decimal import Decimal
from io import BytesIO
from pathlib import Path

import openpyxl
from fastapi.testclient import TestClient

from hesabdar.charts.spec import CHART_TYPES, assert_charts, build_chart, read_chart_title
from hesabdar.files.workbook import write_workbook
from hesabdar.planner.rules import build_plan
from tests.test_operations import _assert_original, _download, _names, _one, _xlsx

TREND_TEXT = "روند فروش ماهانه را نشان بده"
DASHBOARD_TEXT = "یک داشبورد فروش بساز"
FORMULA_TEXT = "فرمول درصد اختلاف را بنویس"
CHART_TEXT = "یک نمودار ستونی از فروش بساز"

SALES = [
    ["شهر", "تاریخ", "مبلغ", "مشتری", "محصول"],
    ["تهران", "1404-07-01", 100, "علی", "مداد"],
    ["شیراز", "1404-07-02", 40, "سارا", "دفتر"],
    ["تهران", "1404-08-01", 50, "علی", "مداد"],
]


def test_packages_do_not_import_each_other() -> None:
    root = Path(__file__).resolve().parents[1] / "hesabdar"
    chart_source = (root / "charts" / "spec.py").read_text()
    dashboard_source = (root / "dashboard" / "layout.py").read_text()
    formula_source = (root / "formulas" / "live.py").read_text()
    assert "hesabdar.analysis" not in chart_source
    assert "hesabdar.reconcile" not in chart_source
    assert "hesabdar.analysis" not in dashboard_source
    assert "hesabdar.charts" not in formula_source


def test_seven_chart_types_round_trip_the_persian_title_and_source_values() -> None:
    table = _table(SALES)
    expected = _city_totals(SALES)
    for chart_type in CHART_TYPES:
        title = f"نمودار {chart_type} فروش"
        category = None if chart_type == "scatter" else "شهر"
        built = build_chart(table, "مبلغ", chart_type, title, category, "فروش")
        assert built["blocked"] is None
        spec = built["spec"]
        if chart_type == "scatter":
            assert [point["value"] for point in spec["series"][0]["points"]] == ["100", "40", "50"]
        else:
            assert [(point["category"], point["value"]) for point in spec["series"][0]["points"]] == expected
            if chart_type == "combo":
                assert [point["value"] for point in spec["series"][1]["points"]] == [value for _city, value in expected]
        data = write_workbook([built["sheet"]], "xlsx")
        assert_charts(data, [("نمودار", spec)])
        book = openpyxl.load_workbook(BytesIO(data))
        assert read_chart_title(book["نمودار"]._charts[0]) == title
        book.close()


def test_trend_sentence_draws_a_chart_from_the_months(client: TestClient) -> None:
    plan = build_plan(TREND_TEXT, "w", ["f"])
    assert [step["tool"] for step in plan["steps"]] == ["inspect", "trend", "chart"]
    payload = _xlsx(SALES)
    workspace, stored = _one(client, "روند", "فروش.xlsx", payload)
    response = client.post(
        f"/api/workspaces/{workspace['id']}/operations",
        json={"text": TREND_TEXT, "file_ids": [stored["id"]]},
    )
    assert response.status_code == 201, response.text
    body = response.json()
    assert body["status"] == "succeeded"
    assert body["error"] is None
    chart = body["charts"][0]
    assert chart["title"] == "روند فروش ماهانه"
    assert chart["chart_type"] == "line"
    assert [(point["category"], point["value"]) for point in chart["series"][0]["points"]] == [("1404-07", "140"), ("1404-08", "50")]
    assert "99" not in body["summary"]
    output = next(item for item in body["outputs"] if item["name"] == "نمودار.xlsx")
    saved = _download(client, workspace["id"], output["version_id"])
    book = openpyxl.load_workbook(BytesIO(saved))
    assert read_chart_title(book["نمودار"]._charts[0]) == "روند فروش ماهانه"
    values = [row[1] for row in book["نمودار"].iter_rows(min_row=2, values_only=True)]
    assert values == [140, 50]
    book.close()
    _assert_original(client, workspace["id"], stored, payload)


def test_dashboard_keeps_only_cards_with_real_columns(client: TestClient) -> None:
    payload = _xlsx(SALES)
    workspace, stored = _one(client, "داشبورد", "فروش.xlsx", payload)
    response = client.post(
        f"/api/workspaces/{workspace['id']}/operations",
        json={"text": DASHBOARD_TEXT, "file_ids": [stored["id"]]},
    )
    assert response.status_code == 201, response.text
    body = response.json()
    assert body["status"] == "succeeded"
    assert _card(body, "total") == "190"
    assert _card(body, "mean") == _plain(Decimal(190) / Decimal(3))
    assert _card(body, "customer_count") == "2"
    assert _card(body, "top_product") == "مداد"
    assert _card(body, "top_city") == "تهران"
    assert all(card["name"] != "category_total" for card in body["cards"])
    assert "دسته‌بندی" in "".join(body["warnings"])
    assert body["charts"][0]["title"] == "روند ماهانه"
    assert [point["value"] for point in body["charts"][0]["series"][0]["points"]] == ["140", "50"]
    saved = _download(client, workspace["id"], body["outputs"][0]["version_id"])
    book = openpyxl.load_workbook(BytesIO(saved))
    assert "محدودیت‌ها" in book.sheetnames
    limits = list(book["محدودیت‌ها"].iter_rows(values_only=True))
    assert ("فروش دسته‌بندی", "قابل محاسبه نیست", "دسته‌بندی") in limits
    products = list(book["محصولات"].iter_rows(values_only=True))
    assert ("مداد", 150) in products
    book.close()
    _assert_original(client, workspace["id"], stored, payload)


def test_dashboard_without_amount_invents_no_card(client: TestClient) -> None:
    payload = _xlsx([["نام", "شهر"], ["علی", "تهران"]])
    workspace, stored = _one(client, "بدون مبلغ", "فروش.xlsx", payload)
    before = _names(client, workspace["id"])
    response = client.post(
        f"/api/workspaces/{workspace['id']}/operations",
        json={"text": DASHBOARD_TEXT, "file_ids": [stored["id"]]},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["status"] == "waiting_confirmation"
    assert body["cards"] == []
    assert body["charts"] == []
    assert body["outputs"] == []
    assert "مبلغ" in body["summary"]
    assert _names(client, workspace["id"]) == before


def test_chart_without_a_category_writes_nothing(client: TestClient) -> None:
    payload = _xlsx([["نام", "مبلغ"], ["علی", 10]])
    workspace, stored = _one(client, "بدون دسته", "فروش.xlsx", payload)
    before = _names(client, workspace["id"])
    response = client.post(
        f"/api/workspaces/{workspace['id']}/operations",
        json={"text": CHART_TEXT, "file_ids": [stored["id"]]},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["status"] == "waiting_confirmation"
    assert body["outputs"] == []
    assert body["charts"] == []
    assert "ستون دسته" in body["summary"]
    assert _names(client, workspace["id"]) == before


def test_live_formula_keeps_the_reference_and_is_not_an_excel_result(client: TestClient) -> None:
    rows = [["نام", "مبلغ", "مبلغ فروش"], ["علی", 100, 110], ["سارا", 0, 20]]
    payload = _xlsx(rows)
    workspace, stored = _one(client, "فرمول", "فروش.xlsx", payload)
    response = client.post(
        f"/api/workspaces/{workspace['id']}/operations",
        json={"text": FORMULA_TEXT, "file_ids": [stored["id"]]},
    )
    assert response.status_code == 201, response.text
    body = response.json()
    assert body["status"] == "succeeded"
    assert body["evaluation"] == "unchecked"
    assert "ارزیابی نشده" in body["summary"]
    assert "نتیجه اکسل" not in body["summary"]
    saved = _download(client, workspace["id"], body["outputs"][0]["version_id"])
    book = openpyxl.load_workbook(BytesIO(saved))
    assert book.active.cell(2, 4).value == '=IF(B2=0,"",ABS(C2-B2)/ABS(B2)*100)'
    assert book.active.cell(3, 4).value == '=IF(B3=0,"",ABS(C3-B3)/ABS(B3)*100)'
    assert ("فرمول زنده", "ارزیابی نشده") in list(book["درباره این خروجی"].iter_rows(values_only=True))
    book.close()
    cached = openpyxl.load_workbook(BytesIO(saved), data_only=True)
    assert cached.active.cell(2, 4).value is None
    cached.close()
    _assert_original(client, workspace["id"], stored, payload)


def _table(rows: list[list]) -> dict:
    headers = rows[0]
    return {
        "sheet": "فروش",
        "headers": headers,
        "rows": [dict(zip(headers, row, strict=True)) for row in rows[1:]],
        "roles": {"شهر": "city", "تاریخ": "date", "مبلغ": "amount", "مشتری": "person_name", "محصول": None},
    }


def _city_totals(rows: list[list]) -> list[tuple[str, str]]:
    totals: dict[str, Decimal] = {}
    order: list[str] = []
    for row in rows[1:]:
        city = row[0]
        if city not in totals:
            order.append(city)
            totals[city] = Decimal(0)
        totals[city] += Decimal(row[2])
    return [(city, _plain(totals[city])) for city in order]


def _plain(value: Decimal) -> str:
    rendered = format(value, "f")
    if "." in rendered:
        rendered = rendered.rstrip("0").rstrip(".")
    return rendered


def _card(body: dict, name: str) -> str:
    return next(card["value"] for card in body["cards"] if card["name"] == name)
