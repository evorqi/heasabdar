from __future__ import annotations

import csv
import hashlib
import zipfile
from io import BytesIO, StringIO
from pathlib import Path

import openpyxl
from fastapi.testclient import TestClient

from hesabdar.messages import CONFIRM_REWRITE, PLAN_NEED_TWO_SHEETS, PLAN_ONE_SHEET
from hesabdar.structure.engine import block_reason, build_management
from tests.test_api import create_workspace, upload

CREATE_TEXT = "یک فایل مدیریت فروش با برگه‌های فروش، مشتریان، محصولات، گزارش و داشبورد بساز"
MERGE_TEXT = "دو برگه را با هم ادغام کن"


def test_structure_package_does_not_import_sibling_engines() -> None:
    root = Path(__file__).resolve().parents[1] / "hesabdar" / "structure"
    source = "\n".join(path.read_text(encoding="utf-8") for path in root.glob("*.py"))
    for name in ("cleaning", "reconcile", "analysis", "charts"):
        assert f"hesabdar.{name}" not in source


def test_management_template_lists_only_named_sheets() -> None:
    full = build_management(CREATE_TEXT)
    assert [sheet["name"] for sheet in full["sheets"]] == ["فروش", "مشتریان", "محصولات", "گزارش", "داشبورد"]
    assert "درباره" not in full["summary"]
    partial = build_management("یک فایل مدیریت فروش بساز")
    assert [sheet["name"] for sheet in partial["sheets"]] == ["فروش"]
    assert "داشبورد" not in partial["summary"]
    assert block_reason(["macro"], False) == "macro"
    assert block_reason(["pivot_table"], False) == "pivot"
    assert block_reason(["macro", "pivot_table"], True) is None


def test_create_workbook_opens_with_the_promised_sheets(client: TestClient) -> None:
    workspace = create_workspace(client, "ساخت فایل")
    response = client.post(
        f"/api/workspaces/{workspace['id']}/operations",
        json={"text": CREATE_TEXT, "file_ids": []},
    )
    assert response.status_code == 201, response.text
    body = response.json()
    assert body["plan"]["goal"] == "create_workbook"
    assert body["error"] is None
    assert "درباره" not in body["summary"]
    for name in ("فروش", "مشتریان", "محصولات", "گزارش", "داشبورد"):
        assert name in body["summary"]
    book = openpyxl.load_workbook(BytesIO(_download(client, workspace["id"], body["outputs"][0]["version_id"])))
    assert book.sheetnames == ["فروش", "مشتریان", "محصولات", "گزارش", "داشبورد"]
    sales = book["فروش"]
    assert sales.sheet_view.rightToLeft is True
    assert sales.freeze_panes == "A2"
    assert sales.cell(1, 1).font.bold is True
    assert sales.cell(1, 1).font.name == "Tahoma"
    assert [cell.value for cell in sales[1]] == ["شماره فاکتور", "تاریخ", "مشتری", "محصول", "شهر", "مبلغ"]
    assert sales.max_row == 1
    assert sales.column_dimensions["A"].number_format == "@"
    assert book["مشتریان"].column_dimensions["A"].number_format == "@"
    assert book["محصولات"].column_dimensions["A"].number_format == "@"
    assert list(book["گزارش"].iter_rows(values_only=True)) == [("شاخص", "مقدار", "واحد")]
    dashboard = book["داشبورد"]
    assert list(dashboard.iter_rows(values_only=True)) == [
        ("موضوع", "وضعیت"),
        ("نمودار", "تا وارد شدن داده فروش ساخته نمی‌شود"),
    ]
    assert dashboard._charts == []
    assert all(not worksheet._charts for worksheet in book.worksheets)
    rules = list(sales.data_validations.dataValidation)
    assert rules and "تهران" in (rules[0].formula1 or "")
    book.close()


def test_sheet_merge_row_count_and_formula_pattern(client: TestClient) -> None:
    payload = _formula_book()
    left_rows, right_rows = _data_counts(payload)
    assert left_rows + right_rows == 4
    workspace, stored = _one(client, "ادغام", "فروش.xlsx", payload)
    response = client.post(
        f"/api/workspaces/{workspace['id']}/operations",
        json={"text": MERGE_TEXT, "file_ids": [stored["id"]]},
    )
    assert response.status_code == 201, response.text
    body = response.json()
    assert body["plan"]["goal"] == "sheet_merge"
    assert body["summary"].startswith(f"{left_rows + right_rows} ردیف")
    warning = " ".join(body["warnings"])
    assert "1" in warning
    assert "=" not in warning
    assert "B2" not in warning
    output = _download(client, workspace["id"], body["outputs"][0]["version_id"])
    book = openpyxl.load_workbook(BytesIO(output))
    sheet = book.active
    values = list(sheet.iter_rows(values_only=True))
    assert len(values) - 1 == left_rows + right_rows
    invoice = sheet.cell(2, 1)
    assert invoice.value == "00123"
    assert invoice.data_type == "s"
    formula = sheet.cell(2, 3)
    assert formula.data_type == "s"
    assert str(formula.value).startswith("=")
    book.close()
    _assert_original(client, workspace["id"], stored, payload)


def test_macro_resave_waits_and_keeps_the_original(client: TestClient) -> None:
    payload = _with_part(_formula_book(), "xl/vbaProject.bin", b"macro")
    workspace, stored = _one(client, "ماکرو", "ماکرو.xlsx", payload)
    response = client.post(
        f"/api/workspaces/{workspace['id']}/operations",
        json={"text": MERGE_TEXT, "file_ids": [stored["id"]]},
    )
    assert response.status_code == 200, response.text
    waiting = response.json()
    assert waiting["status"] == "waiting_confirmation"
    assert waiting["outputs"] == []
    assert "ماکرو" in waiting["summary"]
    confirmed = client.post(
        f"/api/workspaces/{workspace['id']}/operations/{waiting['id']}/confirm",
        json={"column": CONFIRM_REWRITE, "key_columns": []},
    )
    assert confirmed.status_code == 201, confirmed.text
    done = confirmed.json()
    assert done["status"] == "succeeded"
    assert any("ماکرو" in warning for warning in done["warnings"])
    assert "=" not in " ".join(done["warnings"])
    _assert_original(client, workspace["id"], stored, payload)


def test_pivot_resave_waits_without_writing(client: TestClient) -> None:
    payload = _with_part(_plain_book(), "xl/pivotTables/pivotTable1.xml", b"<pivot/>")
    workspace, stored = _one(client, "محوری", "محوری.xlsx", payload)
    response = client.post(
        f"/api/workspaces/{workspace['id']}/operations",
        json={"text": "سطر عنوان را تثبیت کن", "file_ids": [stored["id"]]},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["outputs"] == []
    assert "محوری" in body["summary"]
    _assert_original(client, workspace["id"], stored, payload)


def test_csv_export_neutralizes_dangerous_cells(client: TestClient) -> None:
    rows = [
        ["شماره فاکتور", "شرح"],
        ["00123", "=cmd|'/c calc'!A0"],
        ["00456", "+1"],
        ["00007", "@sum"],
    ]
    payload = _sheet("فروش", rows)
    workspace, stored = _one(client, "csv", "فروش.xlsx", payload)
    response = client.post(
        f"/api/workspaces/{workspace['id']}/operations",
        json={"text": "خروجی csv بده", "file_ids": [stored["id"]]},
    )
    assert response.status_code == 201, response.text
    body = response.json()
    assert body["plan"]["goal"] == "export"
    assert body["outputs"][0]["name"] == "خروجی.csv"
    warning = " ".join(body["warnings"])
    assert warning.startswith("3 ")
    assert "calc" not in body["summary"]
    assert "calc" not in warning
    data = _download(client, workspace["id"], body["outputs"][0]["version_id"])
    text = data.decode("utf-8-sig")
    parsed = list(csv.reader(StringIO(text)))
    assert parsed[1][0] == "00123"
    assert parsed[1][1].startswith("'=")
    assert parsed[2][1].startswith("'+")
    assert parsed[3][1].startswith("'@")
    _assert_original(client, workspace["id"], stored, payload)


def test_sheet_tools_write_or_ask(client: TestClient) -> None:
    payload = _plain_book()
    workspace, stored = _one(client, "ساختار", "جدول.xlsx", payload)
    added = _run(client, workspace["id"], "برگه تازه را اضافه کن", stored["id"])
    book = _book(client, workspace["id"], added)
    assert "تازه" in book.sheetnames
    book.close()
    renamed = _run(client, workspace["id"], "نام برگه فروش را خلاصه بگذار", stored["id"])
    book = _book(client, workspace["id"], renamed)
    assert "خلاصه" in book.sheetnames
    book.close()
    moved = _run(client, workspace["id"], "برگه دوم را اول بگذار", stored["id"])
    book = _book(client, workspace["id"], moved)
    assert book.sheetnames[0] == "دوم"
    book.close()
    deleted = _run(client, workspace["id"], "برگه دوم را حذف کن", stored["id"])
    book = _book(client, workspace["id"], deleted)
    assert "دوم" not in book.sheetnames
    assert len(book.sheetnames) == 1
    book.close()
    single = upload(client, workspace["id"], [("تک.xlsx", _sheet("فروش", [["شهر", "مبلغ"], ["تهران", 1]]))])
    assert single.status_code == 201, single.text
    single_id = single.json()["accepted"][0]["id"]
    alone = client.post(
        f"/api/workspaces/{workspace['id']}/operations",
        json={"text": "برگه فروش را حذف کن", "file_ids": [single_id]},
    )
    assert alone.status_code == 200, alone.text
    assert alone.json()["outputs"] == []
    assert alone.json()["summary"] == PLAN_ONE_SHEET
    missing = client.post(
        f"/api/workspaces/{workspace['id']}/operations",
        json={"text": MERGE_TEXT, "file_ids": [single_id]},
    )
    assert missing.json()["outputs"] == []
    assert missing.json()["summary"] == PLAN_NEED_TWO_SHEETS


def test_column_row_freeze_group_and_split(client: TestClient) -> None:
    rows = [
        ["شهر", "مبلغ", "شرح"],
        ["تهران", 10, "الف"],
        ["شیراز", 20, "ب"],
        ["مشهد", 30, "ج"],
        ["یزد", 40, "د"],
    ]
    payload = _sheet("فروش", rows)
    workspace, stored = _one(client, "ردیف", "فروش.xlsx", payload)
    fresh = _run(client, workspace["id"], "ستون تازه بساز", stored["id"])
    book = _book(client, workspace["id"], fresh)
    assert book.active.cell(1, 4).value == "تازه"
    book.close()
    moved = _run(client, workspace["id"], "ستون مبلغ را اول بگذار", stored["id"])
    book = _book(client, workspace["id"], moved)
    assert book.active.cell(1, 1).value == "مبلغ"
    book.close()
    removed = _run(client, workspace["id"], "ستون شرح را حذف کن", stored["id"])
    book = _book(client, workspace["id"], removed)
    headers = [cell.value for cell in book.active[1]]
    assert "شرح" not in headers
    book.close()
    longer = _run(client, workspace["id"], "یک ردیف اضافه کن", stored["id"])
    book = _book(client, workspace["id"], longer)
    assert book.active.max_row == 6
    book.close()
    shorter = _run(client, workspace["id"], "ردیف را حذف کن", stored["id"])
    book = _book(client, workspace["id"], shorter)
    assert book.active.max_row == 4
    book.close()
    frozen = _run(client, workspace["id"], "سطر عنوان را تثبیت کن", stored["id"])
    book = _book(client, workspace["id"], frozen)
    assert book.active.freeze_panes == "A2"
    book.close()
    bold = _run(client, workspace["id"], "قالب سرستون را پررنگ کن", stored["id"])
    book = _book(client, workspace["id"], bold)
    assert book.active.cell(1, 1).font.bold is True
    book.close()
    grouped = _run(client, workspace["id"], "ردیف‌های داده را گروه‌بندی کن", stored["id"])
    book = _book(client, workspace["id"], grouped)
    assert book.active.row_dimensions[2].outline_level == 1
    book.close()
    choices = _run(client, workspace["id"], "فهرست کشویی ستون شهر را بساز", stored["id"])
    book = _book(client, workspace["id"], choices)
    rules = list(book.active.data_validations.dataValidation)
    assert rules and "تهران" in (rules[0].formula1 or "")
    book.close()
    split = _run(client, workspace["id"], "این برگه را تقسیم کن", stored["id"])
    book = _book(client, workspace["id"], split)
    first = book.worksheets[0].max_row - 1
    second = book.worksheets[1].max_row - 1
    assert first + second == len(rows) - 1
    book.close()
    parts = _run(client, workspace["id"], "این فایل را تقسیم کن", stored["id"])
    assert len(parts["outputs"]) == 2
    counts = []
    for output in parts["outputs"]:
        opened = openpyxl.load_workbook(BytesIO(_download(client, workspace["id"], output["version_id"])))
        counts.append(opened.active.max_row - 1)
        opened.close()
    assert sum(counts) == len(rows) - 1


def _run(client: TestClient, workspace_id: str, text: str, file_id: str) -> dict:
    response = client.post(
        f"/api/workspaces/{workspace_id}/operations",
        json={"text": text, "file_ids": [file_id]},
    )
    assert response.status_code == 201, response.text
    body = response.json()
    assert body["status"] == "succeeded"
    assert body["error"] is None
    assert body["outputs"]
    return body


def _formula_book() -> bytes:
    book = openpyxl.Workbook()
    first = book.active
    first.title = "فروش"
    first.append(["شماره", "مبلغ", "فرمول"])
    first.append(["00123", 10, "=B2*2"])
    first["A2"].number_format = "@"
    first.append(["00124", 20, "=B3*2"])
    first["A3"].number_format = "@"
    first.append(["00125", 30, "=B4+2"])
    first["A4"].number_format = "@"
    second = book.create_sheet("شعبه")
    second.append(["شماره", "مبلغ", "فرمول"])
    second.append(["00126", 5, "=B2*2"])
    second["A2"].number_format = "@"
    return _save(book)


def _plain_book() -> bytes:
    book = openpyxl.Workbook()
    first = book.active
    first.title = "فروش"
    first.append(["شهر", "مبلغ"])
    first.append(["تهران", 10])
    second = book.create_sheet("دوم")
    second.append(["شهر", "مبلغ"])
    second.append(["شیراز", 20])
    return _save(book)


def _sheet(name: str, rows: list[list]) -> bytes:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = name
    for row in rows:
        sheet.append(row)
    if rows and rows[0] and rows[0][0] == "شماره فاکتور":
        for index in range(2, len(rows) + 1):
            sheet.cell(index, 1).number_format = "@"
            sheet.cell(index, 1).value = str(rows[index - 1][0])
    return _save(book)


def _with_part(payload: bytes, name: str, content: bytes) -> bytes:
    buffer = BytesIO()
    with zipfile.ZipFile(BytesIO(payload)) as source, zipfile.ZipFile(buffer, "w") as target:
        for info in source.infolist():
            target.writestr(info, source.read(info.filename))
        target.writestr(name, content)
    return buffer.getvalue()


def _data_counts(payload: bytes) -> tuple[int, int]:
    book = openpyxl.load_workbook(BytesIO(payload))
    counts = tuple(sheet.max_row - 1 for sheet in book.worksheets)
    book.close()
    return counts[0], counts[1]


def _save(book) -> bytes:
    buffer = BytesIO()
    book.save(buffer)
    book.close()
    return buffer.getvalue()


def _one(client: TestClient, title: str, name: str, payload: bytes) -> tuple[dict, dict]:
    workspace = create_workspace(client, title)
    uploaded = upload(client, workspace["id"], [(name, payload)])
    assert uploaded.status_code == 201, uploaded.text
    return workspace, uploaded.json()["accepted"][0]


def _download(client: TestClient, workspace_id: str, version_id: str) -> bytes:
    response = client.get(f"/api/workspaces/{workspace_id}/versions/{version_id}/download")
    assert response.status_code == 200, response.text
    return response.content


def _book(client: TestClient, workspace_id: str, body: dict):
    return openpyxl.load_workbook(BytesIO(_download(client, workspace_id, body["outputs"][0]["version_id"])))


def _assert_original(client: TestClient, workspace_id: str, stored: dict, payload: bytes) -> None:
    downloaded = _download(client, workspace_id, stored["current_version_id"])
    assert downloaded == payload
    assert hashlib.sha256(downloaded).hexdigest() == stored["sha256"]
