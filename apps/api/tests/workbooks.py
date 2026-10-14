from __future__ import annotations

import zipfile
from io import BytesIO
from xml.etree import ElementTree as ET

import xlwt
from openpyxl import Workbook

ERROR_TOKENS = ("#REF!", "#VALUE!", "#N/A", "#DIV/0!", "#NAME?", "#NULL!", "#NUM!")


def sales_workbook() -> bytes:
    book = Workbook()
    sales = book.active
    sales.title = "فروش"
    sales.append(["نام", "کد ملی", "مبلغ", "تاریخ", "محاسبه"])
    sales.append(["سارا احمدی", "0012345678", 1500000, "1403/07/20", "=C2*2"])
    sales.append(["رضا", "0123456789", 2500, "1403/13/01", "#DIV/0!"])
    sales["E3"].data_type = "e"
    sales["B2"].number_format = "@"
    sales["B3"].number_format = "@"

    costs = book.create_sheet("هزینه‌ها")
    costs.append(["شرح", "مبلغ"])
    costs.append(["اجاره", 900000])

    book.create_sheet("خالی")
    return _save(book)


def error_workbook() -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "خطاها"
    sheet.append(["پیام"])
    for token in ERROR_TOKENS:
        sheet.append([token])
        sheet.cell(sheet.max_row, 1).data_type = "e"
    sheet["B2"] = "=A2+1"
    return _save(book)


def flagged_workbook() -> bytes:
    base = sales_workbook()
    return _add_zip_entries(
        base,
        {
            "xl/vbaProject.bin": b"not-a-macro",
            "xl/pivotTables/pivotTable1.xml": b"<pivot/>",
        },
    )


def style_index_workbook() -> bytes:
    """شماره سبکی که در فهرست سبک‌ها نیست. مقدار سلول باید بماند."""
    return _patch_sheet(sales_sheet(), lambda text: text.replace("<c ", '<c s="99" ', 1))


def broken_row_workbook() -> bytes:
    """یک ردیف با شماره نامعتبر، و ردیف‌های سالم بعد از آن."""
    broken = '<row r="1.5"><c r="A1" t="inlineStr"><is><t>بد</t></is></c></row>'
    return _patch_sheet(sales_sheet(), lambda text: text.replace("<sheetData>", f"<sheetData>{broken}", 1))


def sales_sheet() -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "فروش"
    sheet.append(["نام", "مبلغ"])
    sheet.append(["علی", 15000000])
    sheet["C2"] = "=B2*2"
    return _save(book)


def pivot_workbook() -> bytes:
    """برگه فروش سالم، به‌علاوه رابطه جدول محوری که موتور آن را رد می‌کند."""
    data = sales_sheet()
    pivot = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<pivotTableDefinition xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        'name="Pivot1" cacheId="" dataCaption="Values"/>'
    ).encode()
    return _link_part(
        data,
        "xl/pivotTables/pivotTable1.xml",
        pivot,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/pivotTable",
        "../pivotTables/pivotTable1.xml",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.pivotTable+xml",
    )


def persian_csv() -> bytes:
    return "نام,کد ملی,مبلغ\nعلی‌رضا,۰۰۱۲۳۴۵۶۷۸,۱۲۳۴\n".encode("utf-8-sig")


def persian_xls() -> bytes:
    book = xlwt.Workbook(encoding="utf-8")
    sheet = book.add_sheet("فروش")
    sheet.write(0, 0, "نام")
    sheet.write(0, 1, "کد ملی")
    sheet.write(0, 2, "مبلغ")
    sheet.write(1, 0, "سارا")
    sheet.write(1, 1, "0012345678")
    sheet.write(1, 2, 1500)
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _patch_sheet(data: bytes, patch) -> bytes:
    source = zipfile.ZipFile(BytesIO(data))
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for info in source.infolist():
            body = source.read(info.filename)
            if info.filename.endswith("sheet1.xml"):
                body = patch(body.decode()).encode()
            archive.writestr(info.filename, body)
    source.close()
    return buffer.getvalue()


def _save(book: Workbook) -> bytes:
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _link_part(
    data: bytes,
    part_name: str,
    payload: bytes,
    rel_type: str,
    target: str,
    content_type: str,
) -> bytes:
    source = zipfile.ZipFile(BytesIO(data))
    names = {info.filename for info in source.infolist()}
    rels_name = "xl/worksheets/_rels/sheet1.xml.rels"
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for info in source.infolist():
            body = source.read(info.filename)
            if info.filename == "[Content_Types].xml":
                root = ET.fromstring(body)
                ET.SubElement(
                    root,
                    "{http://schemas.openxmlformats.org/package/2006/content-types}Override",
                    PartName=f"/{part_name}",
                    ContentType=content_type,
                )
                body = ET.tostring(root)
            if info.filename == rels_name:
                root = ET.fromstring(body)
                ET.SubElement(
                    root,
                    "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
                    Id="rId9",
                    Type=rel_type,
                    Target=target,
                )
                body = ET.tostring(root)
            archive.writestr(info.filename, body)
        archive.writestr(part_name, payload)
        if rels_name not in names:
            root = ET.Element("{http://schemas.openxmlformats.org/package/2006/relationships}Relationships")
            ET.SubElement(
                root,
                "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
                Id="rId9",
                Type=rel_type,
                Target=target,
            )
            archive.writestr(rels_name, ET.tostring(root))
    source.close()
    return buffer.getvalue()


def _add_zip_entries(data: bytes, entries: dict[str, bytes]) -> bytes:
    source = zipfile.ZipFile(BytesIO(data))
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for info in source.infolist():
            archive.writestr(info, source.read(info.filename))
        for name, payload in entries.items():
            archive.writestr(name, payload)
    source.close()
    return buffer.getvalue()
