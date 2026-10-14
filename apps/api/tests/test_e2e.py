from __future__ import annotations

import hashlib
from io import BytesIO

import openpyxl
from fastapi.testclient import TestClient

from tests.test_api import create_workspace, upload
from tests.test_operations import RECONCILE_TEXT, _download, _xlsx

SENTENCE = "سلام، این یک متن فارسی معمولی است."
MIXED = "فروش شعبه تهران در مهر برابر با 12500000 ریال (IRR) بود."
DIGITS = "۱۲۳۴۵۶"
ARABIC_NAME = "علي كريمي"
ZWNJ_NAME = "علی\u200cرضا"
PHONE = "09121234567"
IDENTIFIER = "0012345678"
EMAIL_SENTENCE = "پاسخ را به user@example.com بفرستید."
PERCENT = "رشد ۱۲٫۵٪ نسبت به مهر"
JALALI = "۱۴۰۳/۰۷/۲۱"
ISO_DATE = "2024-10-12"
AMBIGUOUS_DATE = "02/03/04"
LONG = (
    "این پاراگراف برای آزمون فارسی است. فروش شعبه تهران ثبت شد. "
    "شناسه و تاریخ در خانه جدا می‌مانند و ترتیب نویسه‌ها عوض نمی‌شود."
)


def test_mother_reconcile_flow_keeps_count_and_persian(client: TestClient) -> None:
    left = _xlsx(
        [
            ["شماره فاکتور", "نام", "مبلغ"],
            ["1001", "علی", 1000],
            ["1002", "سارا", 0],
            ["1003", "رضا", 500],
            [None, "نیکان", 200],
        ]
    )
    right = _xlsx(
        [
            ["شماره فاکتور", "نام", "مبلغ"],
            ["۱۰۰۱", "علي", 1000],
            ["1002", "سارا", 50],
            ["1004", "مریم", 300],
            [None, "نيكان", 200],
        ]
    )
    workspace = create_workspace(client, "جریان مادر")
    uploaded = upload(client, workspace["id"], [("فروش اول.xlsx", left), ("فروش دوم.xlsx", right)])
    assert uploaded.status_code == 201, uploaded.text
    accepted = {item["original_name"]: item for item in uploaded.json()["accepted"]}

    preview = client.post(
        f"/api/workspaces/{workspace['id']}/reconciliations/preview",
        json={"left_file_id": accepted["فروش اول.xlsx"]["id"], "right_file_id": accepted["فروش دوم.xlsx"]["id"]},
    )
    assert preview.status_code == 200, preview.text
    assert preview.json()["key_columns"] == ["شماره فاکتور"]
    assert preview.json()["key_ambiguous"] is False
    assert preview.json()["result_file_id"] is None

    operation = client.post(
        f"/api/workspaces/{workspace['id']}/operations",
        json={
            "text": RECONCILE_TEXT,
            "file_ids": [accepted["فروش اول.xlsx"]["id"], accepted["فروش دوم.xlsx"]["id"]],
        },
    )
    assert operation.status_code == 201, operation.text
    body = operation.json()
    assert body["status"] == "succeeded"
    assert body["error"] is None
    assert body["plan"]["goal"] == "reconcile"
    assert body["summary"].startswith("منطبق 1")

    output = _download(client, workspace["id"], body["outputs"][0]["version_id"])
    book = openpyxl.load_workbook(BytesIO(output))
    assert book.sheetnames == ["خلاصه", "منطبق", "مغایر", "فقط فایل اول", "فقط فایل دوم", "مشکوک", "درباره این خروجی"]
    assert all(sheet.sheet_view.rightToLeft for sheet in book.worksheets)
    summary = {book["خلاصه"][f"A{row}"].value: book["خلاصه"][f"B{row}"].value for row in range(2, 13)}
    assert summary["کلید"] == "شماره فاکتور"
    assert summary["تصمیم کلید"] == "detected"
    assert int(summary["منطبق"]) + int(summary["مغایر"]) + int(summary["فقط فایل اول"]) + int(summary["مشکوک"]) == int(
        summary["ردیف چپ"]
    )
    assert int(summary["منطبق"]) + int(summary["مغایر"]) + int(summary["فقط فایل دوم"]) + int(summary["مشکوک"]) == int(
        summary["ردیف راست"]
    )
    assert summary["ردیف چپ"] == "4"
    assert summary["ردیف راست"] == "4"
    assert book["درباره این خروجی"]["B4"].value == "هر ردیف ورودی دقیقاً در یک دسته است."
    matched = _values(book["منطبق"])
    assert "۱۰۰۱" in matched
    assert "1001" in matched
    suspect = _values(book["مشکوک"])
    assert "نیکان" in suspect
    assert "نيكان" in suspect
    book.close()

    assert _download(client, workspace["id"], accepted["فروش اول.xlsx"]["current_version_id"]) == left
    assert _download(client, workspace["id"], accepted["فروش دوم.xlsx"]["current_version_id"]) == right
    listed = client.get(f"/api/workspaces/{workspace['id']}/files")
    stored = {item["original_name"]: item for item in listed.json()["files"]}
    assert stored["فروش اول.xlsx"]["sha256"] == hashlib.sha256(left).hexdigest()
    assert stored["فروش دوم.xlsx"]["sha256"] == hashlib.sha256(right).hexdigest()


def test_persian_file_roundtrip_and_csv(client: TestClient) -> None:
    rows = [
        ["نام مشتری", "مبلغ فروش (ریال)", "درصد اختلاف", "یادداشت"],
        [SENTENCE, DIGITS, PERCENT, EMAIL_SENTENCE],
        [ARABIC_NAME, IDENTIFIER, PHONE, ZWNJ_NAME],
        [MIXED, JALALI, ISO_DATE, AMBIGUOUS_DATE],
        [LONG, "12500000", "12.5%", "user@example.com"],
    ]
    payload = _xlsx(rows)
    workspace = create_workspace(client, "فارسی")
    uploaded = upload(client, workspace["id"], [("نمونه.xlsx", payload)])
    assert uploaded.status_code == 201, uploaded.text
    stored = uploaded.json()["accepted"][0]
    report = client.post(f"/api/workspaces/{workspace['id']}/inspections", json={"file_id": stored["id"]})
    assert report.status_code == 200, report.text
    raws = _raws(report.json())
    for text in (
        SENTENCE,
        MIXED,
        DIGITS,
        ARABIC_NAME,
        ZWNJ_NAME,
        PHONE,
        IDENTIFIER,
        EMAIL_SENTENCE,
        PERCENT,
        JALALI,
        ISO_DATE,
        AMBIGUOUS_DATE,
        LONG,
        "user@example.com",
    ):
        assert text in raws
    assert "علی کریمی" not in raws

    book = openpyxl.load_workbook(BytesIO(_download(client, workspace["id"], stored["current_version_id"])))
    assert book.active["A2"].value == SENTENCE
    assert book.active["B3"].value == IDENTIFIER
    assert book.active["D3"].value == ZWNJ_NAME
    assert "\u200c" in book.active["D3"].value
    assert book.active["D4"].value == AMBIGUOUS_DATE
    book.close()

    sorted_file = client.post(
        f"/api/workspaces/{workspace['id']}/operations",
        json={"text": "داده‌ها را بر اساس مبلغ مرتب کن", "file_ids": [stored["id"]]},
    )
    assert sorted_file.status_code == 201, sorted_file.text
    written = openpyxl.load_workbook(BytesIO(_download(client, workspace["id"], sorted_file.json()["outputs"][0]["version_id"])))
    assert written.active.sheet_view.rightToLeft is True
    written_values = [cell.value for row in written.active.iter_rows() for cell in row]
    assert SENTENCE in written_values
    assert ARABIC_NAME in written_values
    assert ZWNJ_NAME in written_values
    written.close()

    exported = client.post(
        f"/api/workspaces/{workspace['id']}/operations",
        json={"text": "خروجی csv بده", "file_ids": [stored["id"]]},
    )
    assert exported.status_code == 201, exported.text
    csv_bytes = _download(client, workspace["id"], exported.json()["outputs"][0]["version_id"])
    assert csv_bytes.startswith(b"\xef\xbb\xbf")
    text = csv_bytes.decode("utf-8-sig")
    assert text.split(",", 1)[0] == "نام مشتری"
    assert SENTENCE in text
    assert EMAIL_SENTENCE in text
    assert ZWNJ_NAME in text
    assert ARABIC_NAME in text
    assert IDENTIFIER in text


def _values(sheet) -> list[str]:
    found = []
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                found.append(cell.value)
    return found


def _raws(report: dict) -> list[str]:
    found = []
    for sheet in report["sheets"]:
        for row in sheet["preview"]["rows"]:
            for cell in row:
                if cell.get("raw"):
                    found.append(cell["raw"])
    return found
