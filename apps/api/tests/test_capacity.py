from __future__ import annotations

import os
import time
from base64 import b64encode
from io import BytesIO

import openpyxl
from fastapi.testclient import TestClient
from openpyxl import Workbook

from hesabdar.inspection.readers import MAX_SCAN_ROWS
from hesabdar.messages import SHEET_TOO_LONG
from tests.test_api import create_workspace, upload
from tests.test_operations import FILTER_TEXT, _download


def test_row_limit_message_matches_the_reader() -> None:
    assert str(MAX_SCAN_ROWS) in SHEET_TOO_LONG
    assert "ناقص ساخته نشد" in SHEET_TOO_LONG


def test_sheet_at_the_row_ceiling_filters(client: TestClient) -> None:
    data_rows = MAX_SCAN_ROWS - 1
    payload = _wide_rows(data_rows)
    workspace = create_workspace(client, "سقف ردیف")
    started = time.perf_counter()
    uploaded = upload(client, workspace["id"], [("فروش.xlsx", payload)])
    assert uploaded.status_code == 201, uploaded.text
    stored = uploaded.json()["accepted"][0]
    report = client.post(f"/api/workspaces/{workspace['id']}/inspections", json={"file_id": stored["id"]})
    assert report.status_code == 200, report.text
    body = report.json()
    assert body["sheets"][0]["data_row_count"] == data_rows
    assert all(item["code"] != "sheet.truncated" for item in body["unchecked"])
    assert body["sheets"][0]["preview"]["rows"][0][0]["raw"] == "00000000"
    operation = client.post(
        f"/api/workspaces/{workspace['id']}/operations",
        json={"text": FILTER_TEXT, "file_ids": [stored["id"]]},
    )
    elapsed = time.perf_counter() - started
    assert operation.status_code == 201, operation.text
    result = operation.json()
    assert result["status"] == "succeeded"
    assert result["error"] is None
    assert "25000" in result["summary"]
    assert str(data_rows) in result["summary"]
    assert result["change"]["rows_before"] == data_rows
    assert result["change"]["rows_after"] == 25000
    output = _download(client, workspace["id"], result["outputs"][0]["version_id"])
    book = openpyxl.load_workbook(BytesIO(output))
    identifier = book.active["A2"]
    assert identifier.value == "00000000"
    assert identifier.data_type == "s"
    assert book.active["B2"].value == "علی"
    assert book.active.sheet_view.rightToLeft is True
    book.close()
    assert elapsed < 120


def test_sheet_past_the_row_ceiling_does_not_succeed(client: TestClient) -> None:
    payload = _wide_rows(MAX_SCAN_ROWS)
    workspace = create_workspace(client, "بالای سقف")
    uploaded = upload(client, workspace["id"], [("فروش.xlsx", payload)])
    assert uploaded.status_code == 201, uploaded.text
    stored = uploaded.json()["accepted"][0]
    report = client.post(f"/api/workspaces/{workspace['id']}/inspections", json={"file_id": stored["id"]})
    assert report.status_code == 200, report.text
    assert any(item["code"] == "sheet.truncated" for item in report.json()["unchecked"])
    operation = client.post(
        f"/api/workspaces/{workspace['id']}/operations",
        json={"text": FILTER_TEXT, "file_ids": [stored["id"]]},
    )
    assert operation.status_code == 422, operation.text
    body = operation.json()
    assert body["status"] == "failed"
    assert body["status"] != "succeeded"
    assert body["outputs"] == []
    assert body["summary"] is None
    assert body["error"]["message"] == SHEET_TOO_LONG
    names = {item["original_name"] for item in client.get(f"/api/workspaces/{workspace['id']}/files").json()["files"]}
    assert names == {"فروش.xlsx"}
    assert _download(client, workspace["id"], stored["current_version_id"]) == payload


def test_file_near_the_upload_ceiling_is_read(client: TestClient) -> None:
    payload = _bulky(rows=9500, width=3000)
    assert 20 * 1024 * 1024 < len(payload) < 30 * 1024 * 1024
    workspace = create_workspace(client, "سقف حجم")
    started = time.perf_counter()
    uploaded = upload(client, workspace["id"], [("حجم.xlsx", payload)])
    assert uploaded.status_code == 201, uploaded.text
    stored = uploaded.json()["accepted"][0]
    report = client.post(f"/api/workspaces/{workspace['id']}/inspections", json={"file_id": stored["id"]})
    elapsed = time.perf_counter() - started
    assert report.status_code == 200, report.text
    assert report.json()["sheets"][0]["data_row_count"] == 9500
    assert all(item["code"] != "sheet.truncated" for item in report.json()["unchecked"])
    assert _download(client, workspace["id"], stored["current_version_id"]) == payload
    assert elapsed < 60


def _bulky(rows: int, width: int) -> bytes:
    book = Workbook(write_only=True)
    sheet = book.create_sheet("فروش")
    sheet.append(["یادداشت", "مبلغ"])
    for index in range(rows):
        sheet.append([b64encode(os.urandom(width)).decode("ascii"), 1000 + index])
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _wide_rows(data_rows: int) -> bytes:
    book = Workbook(write_only=True)
    sheet = book.create_sheet("فروش")
    sheet.append(["شماره فاکتور", "نام", "مبلغ", "تاریخ"])
    for index in range(data_rows):
        sheet.append([f"{index:08d}", "علی", 60000000 if index % 2 == 0 else 1000, "1403/07/21"])
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()
