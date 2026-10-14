from __future__ import annotations

import hashlib
import json
import sqlite3
from decimal import Decimal
from io import BytesIO

import openpyxl
import pytest
from fastapi.testclient import TestClient

from hesabdar.config import Settings
from hesabdar.errors import AppError
from hesabdar.planner.explain import consistent
from hesabdar.planner.llm import llm_status
from hesabdar.planner.registry import TOOLS
from hesabdar.planner.rules import build_plan, repeat_period
from hesabdar.planner.validate import validate_plan
from tests.test_api import create_workspace, upload

FILTER_TEXT = "فروش بیشتر از ۵۰ میلیون را پیدا کن"
SORT_TEXT = "داده‌ها را بر اساس مبلغ مرتب کن"
COMPUTE_TEXT = "ستون درصد اختلاف را اضافه کن"
RECONCILE_TEXT = "این دو فایل را تطبیق بده، اختلاف مبالغ را پیدا کن و فایل مغایرت بده"
RANK_TEXT = "بیشترین فروش متعلق به کدام شهر است"
REPEAT_TEXT = "همین گزارش را برای مهر هم بساز"


def test_registry_stops_before_later_tools() -> None:
    assert {"filter", "sort", "add_column", "compute_column", "reconcile", "inspect", "clean"} <= TOOLS
    assert {"aggregate", "rank", "trend", "outliers", "kpi", "report"} <= TOOLS
    assert {"chart", "dashboard", "apply_formula"} <= TOOLS
    assert {
        "create_workbook",
        "sheet_add",
        "sheet_delete",
        "sheet_rename",
        "sheet_move",
        "sheet_merge",
        "sheet_split",
        "file_merge",
        "file_split",
        "column_add",
        "column_delete",
        "column_move",
        "row_add",
        "row_delete",
        "format_cells",
        "freeze",
        "data_validation",
        "export",
    } <= TOOLS
    assert "shell" not in TOOLS
    assert "restore_version" in TOOLS


def test_benchmark_sentences_reach_their_tools() -> None:
    assert build_plan(FILTER_TEXT, "w", ["f"])["goal"] == "filter"
    assert build_plan(SORT_TEXT, "w", ["f"])["goal"] == "sort"
    assert build_plan(COMPUTE_TEXT, "w", ["f"])["goal"] == "compute_column"
    assert build_plan(RECONCILE_TEXT, "w", ["a", "b"])["goal"] == "reconcile"
    vague = build_plan("سلام، یک کار انجام بده", "w", ["f"])
    assert vague["steps"] == []
    assert vague["ambiguities"][0]["code"] == "plan.ambiguous"
    assert vague["ambiguities"][0]["question"]
    rank = build_plan(RANK_TEXT, "w", ["f"])
    assert rank["goal"] == "rank"
    assert [step["tool"] for step in rank["steps"]] == ["inspect", "rank"]
    report = build_plan("یک گزارش مدیریتی بساز", "w", ["f"])
    assert report["goal"] == "report"
    trend = build_plan("روند فروش ماهانه را نشان بده", "w", ["f"])
    assert trend["goal"] == "trend"
    assert [step["tool"] for step in trend["steps"]] == ["inspect", "trend", "chart"]
    assert trend["steps"][-1]["params"]["chart_type"] == "line"
    dashboard = build_plan("یک داشبورد فروش بساز", "w", ["f"])
    assert dashboard["goal"] == "dashboard"
    assert [step["tool"] for step in dashboard["steps"]] == ["inspect", "dashboard"]
    created = build_plan(
        "یک فایل مدیریت فروش با برگه‌های فروش، مشتریان، محصولات، گزارش و داشبورد بساز",
        "w",
        [],
    )
    assert created["goal"] == "create_workbook"
    assert [step["tool"] for step in created["steps"]] == ["create_workbook"]
    merged = build_plan("دو برگه را با هم ادغام کن", "w", ["f"])
    assert merged["goal"] == "sheet_merge"
    assert [step["tool"] for step in merged["steps"]] == ["inspect", "sheet_merge"]
    files = build_plan("دو فایل را ادغام کن", "w", ["a", "b"])
    assert files["goal"] == "file_merge"
    assert "reconcile" not in [step["tool"] for step in files["steps"]]
    restored = build_plan("به نسخه اصلی برگرد", "w", ["f"])
    assert restored["goal"] == "restore_version"
    assert [step["tool"] for step in restored["steps"]] == ["restore_version"]
    assert restored["steps"][0]["params"]["target"] == "original"


def test_shell_plan_and_foreign_numbers_are_rejected(tmp_path) -> None:
    with pytest.raises(AppError) as caught:
        validate_plan({"steps": [{"tool": "shell", "params": {"code": "print(1)"}, "reads": [], "writes": []}]})
    assert caught.value.code == "plan.unknown_tool"
    with pytest.raises(AppError):
        validate_plan(
            {"steps": [{"tool": "filter", "params": {"path": "/etc/passwd"}, "reads": [], "writes": []}]}
        )
    with pytest.raises(AppError):
        build_plan("import os\nos.system('rm -rf /')", "w", ["f"])
    assert llm_status(Settings(data_dir=tmp_path)) == "پیکربندی نشده"
    metrics = {"matched_rows": "2", "source_rows": "4"}
    assert consistent("2 ردیف از 4 ردیف شرط را داشتند.", metrics)
    assert not consistent("99 ردیف پیدا شد.", metrics)
    assert not consistent("۹۹ ردیف پیدا شد.", metrics)


def test_filter_count_matches_an_independent_count(client: TestClient) -> None:
    rows = [
        ["نام", "مبلغ"],
        ["علی", 60000000],
        ["سارا", 50000000],
        ["رضا", 40000000],
        ["نیکا", 51000000],
        ["فروش بیشتر از ۵۰ میلیون را پیدا کن", 1000],
        ["=cmd|'/c calc'!A0", 70000000],
    ]
    payload = _xlsx(rows)
    workspace, stored = _one(client, "فیلتر فروش", "فروش.xlsx", payload)
    response = client.post(
        f"/api/workspaces/{workspace['id']}/operations",
        json={"text": FILTER_TEXT, "file_ids": [stored["id"]]},
    )
    assert response.status_code == 201, response.text
    body = response.json()
    assert body["status"] == "succeeded"
    assert body["llm_status"] == "پیکربندی نشده"
    assert body["plan"]["goal"] == "filter"
    assert body["plan"]["user_text"] == FILTER_TEXT
    assert body["plan"]["source"] == "rules"
    assert body["error"] is None
    expected = _independent_filter(rows, Decimal("50000000"))
    assert expected == 3
    assert body["summary"].startswith(f"{expected} ردیف از {len(rows) - 1} ردیف")
    assert "99" not in body["summary"]
    output = _download(client, workspace["id"], body["outputs"][0]["version_id"])
    book = openpyxl.load_workbook(BytesIO(output))
    sheet = book.active
    assert sheet.sheet_view.rightToLeft is True
    values = list(sheet.iter_rows(values_only=True))
    assert len(values) - 1 == expected
    names = [row[0] for row in values[1:]]
    assert "فروش بیشتر از ۵۰ میلیون را پیدا کن" not in names
    assert "=cmd|'/c calc'!A0" in names
    command = sheet.cell(names.index("=cmd|'/c calc'!A0") + 2, 1)
    assert command.data_type == "s"
    book.close()
    _assert_original(client, workspace["id"], stored, payload)

    ranked = client.post(
        f"/api/workspaces/{workspace['id']}/operations",
        json={"text": RANK_TEXT, "file_ids": [stored["id"]]},
    )
    assert ranked.status_code == 200, ranked.text
    assert ranked.json()["status"] == "waiting_confirmation"
    assert ranked.json()["outputs"] == []
    assert ranked.json()["plan"]["goal"] == "rank"
    assert "ستون شهر" in ranked.json()["summary"]
    assert all(step["tool"] != "chart" for step in ranked.json()["plan"]["steps"])
    assert "رتبه.xlsx" not in _names(client, workspace["id"])


def test_sort_uses_the_saved_file_and_keeps_every_row(client: TestClient) -> None:
    rows = [
        ["نام", "مبلغ"],
        ["علی", 10],
        ["=cmd|'/c calc'!A0", 30],
        ["سارا", 20],
    ]
    payload = _xlsx(rows)
    workspace, stored = _one(client, "مرتب‌سازی", "فروش.xlsx", payload)
    client.post(
        f"/api/workspaces/{workspace['id']}/operations",
        json={"text": FILTER_TEXT, "file_ids": [stored["id"]]},
    )
    response = client.post(
        f"/api/workspaces/{workspace['id']}/operations",
        json={"text": SORT_TEXT},
    )
    assert response.status_code == 201, response.text
    body = response.json()
    assert body["plan"]["goal"] == "sort"
    assert body["summary"].startswith("3 ردیف")
    output = _download(client, workspace["id"], body["outputs"][0]["version_id"])
    book = openpyxl.load_workbook(BytesIO(output))
    amounts = [row[1] for row in book.active.iter_rows(min_row=2, values_only=True)]
    assert amounts == ["30", "20", "10"]
    assert any(row[0] == "=cmd|'/c calc'!A0" for row in book.active.iter_rows(min_row=2, values_only=True))
    book.close()
    _assert_original(client, workspace["id"], stored, payload)


def test_compute_column_percent_comes_from_decimal(client: TestClient) -> None:
    rows = [
        ["مبلغ اول", "مبلغ دوم"],
        [100, 80],
        [0, 50],
        [200, 100],
    ]
    payload = _xlsx(rows)
    workspace, stored = _one(client, "درصد", "مبالغ.xlsx", payload)
    response = client.post(
        f"/api/workspaces/{workspace['id']}/operations",
        json={"text": COMPUTE_TEXT, "file_ids": [stored["id"]]},
    )
    assert response.status_code == 201, response.text
    body = response.json()
    assert body["plan"]["goal"] == "compute_column"
    assert body["summary"] == "ستون درصد اختلاف برای 2 ردیف محاسبه شد."
    assert any("1 ردیف" in warning for warning in body["warnings"])
    output = _download(client, workspace["id"], body["outputs"][0]["version_id"])
    book = openpyxl.load_workbook(BytesIO(output))
    headers = [cell.value for cell in next(book.active.iter_rows(max_row=1))]
    assert "درصد اختلاف" in headers
    index = headers.index("درصد اختلاف")
    found = [row[index] for row in book.active.iter_rows(min_row=2, values_only=True)]
    assert found == ["20", None, "50"]
    assert found[0] == _percent(Decimal(100), Decimal(80))
    assert found[2] == _percent(Decimal(200), Decimal(100))
    book.close()
    _assert_original(client, workspace["id"], stored, payload)


def test_add_column_and_code_request_write_nothing_dangerous(client: TestClient) -> None:
    payload = _xlsx([["نام", "مبلغ"], ["علی", 10], ["سارا", 20]])
    workspace, stored = _one(client, "ستون", "فروش.xlsx", payload)
    added = client.post(
        f"/api/workspaces/{workspace['id']}/operations",
        json={"text": "ستون یادداشت را اضافه کن", "file_ids": [stored["id"]]},
    )
    assert added.status_code == 201, added.text
    output = _download(client, workspace["id"], added.json()["outputs"][0]["version_id"])
    book = openpyxl.load_workbook(BytesIO(output))
    headers = [cell.value for cell in next(book.active.iter_rows(max_row=1))]
    assert headers[-1] == "یادداشت"
    assert [row[-1] for row in book.active.iter_rows(min_row=2, values_only=True)] == [None, None]
    book.close()
    before = {item["original_name"] for item in client.get(f"/api/workspaces/{workspace['id']}/files").json()["files"]}
    rejected = client.post(
        f"/api/workspaces/{workspace['id']}/operations",
        json={"text": "import os\nos.system('rm -rf /')", "file_ids": [stored["id"]]},
    )
    assert rejected.status_code == 422, rejected.text
    assert rejected.json()["error"]["code"] == "plan.unknown_tool"
    assert rejected.json()["status"] == "failed"
    assert rejected.json()["outputs"] == []
    assert "os.system" not in rejected.json()["error"]["message"]
    assert "rm" not in rejected.json()["error"]["message"]
    after = {item["original_name"] for item in client.get(f"/api/workspaces/{workspace['id']}/files").json()["files"]}
    assert after == before
    _assert_original(client, workspace["id"], stored, payload)


def test_clean_sentence_keeps_the_original_bytes(client: TestClient) -> None:
    rows = [["نام", "مبلغ"], ["  سارا  ", 10], ["علی", 20], ["رضا", 30], ["نیکا", 40], ["مریم", 50], ["هدی", 60]]
    payload = _xlsx(rows)
    workspace, stored = _one(client, "پاک‌سازی", "نام‌ها.xlsx", payload)
    response = client.post(
        f"/api/workspaces/{workspace['id']}/operations",
        json={"text": "این فایل را پاکسازی کن", "file_ids": [stored["id"]]},
    )
    assert response.status_code == 201, response.text
    assert response.json()["plan"]["goal"] == "clean"
    assert response.json()["outputs"]
    _assert_original(client, workspace["id"], stored, payload)
    cleaned = _download(client, workspace["id"], response.json()["outputs"][0]["version_id"])
    book = openpyxl.load_workbook(BytesIO(cleaned))
    assert book.active["A2"].value == "سارا"
    book.close()


def test_reconcile_sentence_uses_the_same_engine(client: TestClient) -> None:
    left = _xlsx([["شماره فاکتور", "مبلغ"], ["1001", 1000], ["1002", 40]])
    right = _xlsx([["شماره فاکتور", "مبلغ"], ["۱۰۰۱", 1000], ["1003", 10]])
    workspace = create_workspace(client, "مغایرت جمله")
    uploaded = upload(client, workspace["id"], [("چپ.xlsx", left), ("راست.xlsx", right)])
    accepted = {item["original_name"]: item for item in uploaded.json()["accepted"]}
    response = client.post(
        f"/api/workspaces/{workspace['id']}/operations",
        json={
            "text": RECONCILE_TEXT,
            "file_ids": [accepted["چپ.xlsx"]["id"], accepted["راست.xlsx"]["id"]],
        },
    )
    assert response.status_code == 201, response.text
    body = response.json()
    assert body["plan"]["steps"][-1]["tool"] == "reconcile"
    assert body["summary"].startswith("منطبق 1")
    output = _download(client, workspace["id"], body["outputs"][0]["version_id"])
    book = openpyxl.load_workbook(BytesIO(output))
    assert book.sheetnames == ["خلاصه", "منطبق", "مغایر", "فقط فایل اول", "فقط فایل دوم", "مشکوک", "درباره این خروجی"]
    assert all(sheet.sheet_view.rightToLeft for sheet in book.worksheets)
    matched = [cell.value for row in book["منطبق"].iter_rows() for cell in row]
    assert "۱۰۰۱" in matched
    book.close()
    _assert_original(client, workspace["id"], accepted["چپ.xlsx"], left)
    _assert_original(client, workspace["id"], accepted["راست.xlsx"], right)


def test_ambiguous_key_and_vague_text_do_not_write(client: TestClient) -> None:
    left = _xlsx([["کد ملی", "شماره فاکتور", "مبلغ"], ["001", "1001", 10], ["002", "1002", 20]])
    right = _xlsx([["کد ملی", "شماره فاکتور", "مبلغ"], ["001", "1001", 10], ["002", "1002", 20]])
    workspace = create_workspace(client, "کلید")
    uploaded = upload(client, workspace["id"], [("چپ.xlsx", left), ("راست.xlsx", right)])
    accepted = {item["original_name"]: item for item in uploaded.json()["accepted"]}
    vague = client.post(
        f"/api/workspaces/{workspace['id']}/operations",
        json={"text": "سلام، یک کار انجام بده", "file_ids": [accepted["چپ.xlsx"]["id"]]},
    )
    assert vague.status_code == 200, vague.text
    assert vague.json()["outputs"] == []
    assert vague.json()["summary"]
    confirmed = client.post(
        f"/api/workspaces/{workspace['id']}/operations/{vague.json()['id']}/confirm",
        json={},
    )
    assert confirmed.status_code == 200
    assert confirmed.json()["outputs"] == []
    waiting = client.post(
        f"/api/workspaces/{workspace['id']}/operations",
        json={
            "text": RECONCILE_TEXT,
            "file_ids": [accepted["چپ.xlsx"]["id"], accepted["راست.xlsx"]["id"]],
        },
    )
    assert waiting.status_code == 200, waiting.text
    assert waiting.json()["status"] == "waiting_confirmation"
    assert waiting.json()["outputs"] == []
    assert set(waiting.json()["plan"]["ambiguities"][0]["choices"]) == {"کد ملی", "شماره فاکتور"}
    names = {item["original_name"] for item in client.get(f"/api/workspaces/{workspace['id']}/files").json()["files"]}
    assert "مغایرت.xlsx" not in names
    chosen = client.post(
        f"/api/workspaces/{workspace['id']}/operations/{waiting.json()['id']}/confirm",
        json={"key_columns": ["شماره فاکتور"]},
    )
    assert chosen.status_code == 201, chosen.text
    assert chosen.json()["summary"].startswith("منطبق 2")
    names = {item["original_name"] for item in client.get(f"/api/workspaces/{workspace['id']}/files").json()["files"]}
    assert "مغایرت.xlsx" in names


def test_repeat_period_without_a_previous_report_asks(client: TestClient) -> None:
    assert repeat_period(REPEAT_TEXT) == "مهر"
    assert repeat_period(FILTER_TEXT) is None
    assert repeat_period(RANK_TEXT) is None
    payload = _xlsx([["نام", "مبلغ"], ["علی", 60000000]])
    workspace, stored = _one(client, "بدون گزارش", "فروش.xlsx", payload)
    before = _names(client, workspace["id"])
    response = client.post(
        f"/api/workspaces/{workspace['id']}/operations",
        json={"text": REPEAT_TEXT, "file_ids": [stored["id"]]},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["status"] == "waiting_confirmation"
    assert body["outputs"] == []
    assert "کدام گزارش" in body["summary"]
    assert "مهر" in body["summary"]
    assert body["plan"]["steps"] == []
    assert _names(client, workspace["id"]) == before
    _assert_original(client, workspace["id"], stored, payload)


def test_repeat_period_replays_the_previous_filter(client: TestClient, settings: Settings) -> None:
    rows = [["نام", "مبلغ"], ["علی", 60000000], ["سارا", 1000]]
    payload = _xlsx(rows)
    workspace, stored = _one(client, "گزارش مهر", "فروش.xlsx", payload)
    first = client.post(
        f"/api/workspaces/{workspace['id']}/operations",
        json={"text": FILTER_TEXT, "file_ids": [stored["id"]]},
    )
    assert first.status_code == 201, first.text
    first_body = first.json()
    second = client.post(
        f"/api/workspaces/{workspace['id']}/operations",
        json={"text": REPEAT_TEXT, "file_ids": [stored["id"]]},
    )
    assert second.status_code == 201, second.text
    body = second.json()
    assert body["id"] != first_body["id"]
    assert body["status"] == "succeeded"
    assert body["summary"].startswith(first_body["summary"])
    assert body["summary"].endswith("این نتیجه برای مهر است.")
    assert "مهر" in body["summary"]
    assert any(step["tool"] == "filter" for step in body["plan"]["steps"])
    assert body["outputs"]
    assert body["outputs"][0]["version_id"] != first_body["outputs"][0]["version_id"]
    names = _names(client, workspace["id"])
    assert names.count("فیلتر.xlsx") == 2
    _assert_original(client, workspace["id"], stored, payload)
    focus = _focus(settings, workspace["id"])
    assert focus["last_plan_id"] == body["id"]
    assert focus["last_period_label"] == "مهر"


def test_cancel_stops_before_an_output_file(client: TestClient) -> None:
    payload = _xlsx([["نام", "مبلغ"], ["علی", 60000000], ["سارا", 1000]])
    workspace, stored = _one(client, "لغو", "فروش.xlsx", payload)
    operations = client.app.state.ctx.operations
    operations.arm_pause()
    try:
        response = client.post(
            f"/api/workspaces/{workspace['id']}/operations",
            json={"text": FILTER_TEXT, "file_ids": [stored["id"]]},
        )
        assert response.status_code == 200, response.text
        body = response.json()
        assert body["status"] == "running"
        assert body["progress_total"] > 0
        assert body["progress_label"]
        assert body["outputs"] == []
        current = client.get(f"/api/workspaces/{workspace['id']}/operations/{body['id']}")
        assert current.status_code == 200
        assert current.json()["status"] == "running"
        again = client.post(
            f"/api/workspaces/{workspace['id']}/operations",
            json={"text": FILTER_TEXT, "file_ids": [stored["id"]]},
        )
        assert again.status_code == 200, again.text
        assert again.json()["id"] == body["id"]
        assert again.json()["outputs"] == []
        cancelled = client.post(f"/api/workspaces/{workspace['id']}/operations/{body['id']}/cancel")
        assert cancelled.status_code == 200, cancelled.text
        done = cancelled.json()
        assert done["status"] == "cancelled"
        assert done["outputs"] == []
        assert done["error"] is None
        assert done["progress_label"] == "لغو شد"
        assert "فیلتر انجام شد" not in (done["summary"] or "")
        assert "ردیف از" not in (done["summary"] or "")
        assert "فیلتر.xlsx" not in _names(client, workspace["id"])
        _assert_original(client, workspace["id"], stored, payload)
    finally:
        operations.release()


def test_cancel_of_finished_work_keeps_the_file(client: TestClient) -> None:
    payload = _xlsx([["نام", "مبلغ"], ["علی", 60000000]])
    workspace, stored = _one(client, "تمام‌شده", "فروش.xlsx", payload)
    created = client.post(
        f"/api/workspaces/{workspace['id']}/operations",
        json={"text": FILTER_TEXT, "file_ids": [stored["id"]]},
    )
    assert created.status_code == 201, created.text
    body = created.json()
    cancelled = client.post(f"/api/workspaces/{workspace['id']}/operations/{body['id']}/cancel")
    assert cancelled.status_code == 201, cancelled.text
    assert cancelled.json()["status"] == "succeeded"
    assert cancelled.json()["outputs"] == body["outputs"]
    assert "فیلتر.xlsx" in _names(client, workspace["id"])


def _names(client: TestClient, workspace_id: str) -> list[str]:
    listed = client.get(f"/api/workspaces/{workspace_id}/files")
    assert listed.status_code == 200, listed.text
    return [item["original_name"] for item in listed.json()["files"]]


def _focus(settings: Settings, workspace_id: str) -> dict:
    connection = sqlite3.connect(settings.data_dir / "hesabdar.sqlite")
    row = connection.execute(
        "SELECT focus_json FROM conversation_focus WHERE workspace_id = ?",
        (workspace_id,),
    ).fetchone()
    connection.close()
    assert row is not None
    return json.loads(row[0])


def _independent_filter(rows: list[list], threshold: Decimal) -> int:
    kept = 0
    for row in rows[1:]:
        amount = Decimal(row[1])
        if amount > threshold:
            kept += 1
    return kept


def _percent(base: Decimal, other: Decimal) -> str:
    value = (abs(other - base) / abs(base)) * Decimal(100)
    rendered = format(value.quantize(Decimal("0.0001")), "f")
    if "." in rendered:
        rendered = rendered.rstrip("0").rstrip(".")
    return rendered


def _xlsx(rows: list[list]) -> bytes:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "فروش"
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    book.save(buffer)
    book.close()
    return buffer.getvalue()


def _one(client: TestClient, title: str, name: str, payload: bytes) -> tuple[dict, dict]:
    workspace = create_workspace(client, title)
    uploaded = upload(client, workspace["id"], [(name, payload)])
    assert uploaded.status_code == 201, uploaded.text
    return workspace, uploaded.json()["accepted"][0]


def _download(client: TestClient, workspace_id: str, version_id: str) -> bytes:
    response = client.get(f"/api/workspaces/{workspace_id}/versions/{version_id}/download")
    assert response.status_code == 200, response.text
    return response.content


def _assert_original(client: TestClient, workspace_id: str, stored: dict, payload: bytes) -> None:
    downloaded = _download(client, workspace_id, stored["current_version_id"])
    assert downloaded == payload
    assert hashlib.sha256(downloaded).hexdigest() == stored["sha256"]
