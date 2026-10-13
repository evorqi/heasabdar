from __future__ import annotations

from decimal import Decimal
from io import BytesIO

import openpyxl
import pytest
from fastapi.testclient import TestClient

from hesabdar.analysis.engine import analyze, present
from hesabdar.planner.explain import consistent, render
from hesabdar.planner.rules import build_plan
from tests.test_operations import _assert_original, _download, _names, _one, _xlsx

REPORT_TEXT = "یک گزارش مدیریتی بساز"
RANK_TEXT = "بیشترین فروش متعلق به کدام شهر است"
FULL_ROWS = [
    ["شهر", "تاریخ", "مبلغ"],
    ["تهران", "1404-07-01", 10],
    ["تهران", "1404-07-02", 20],
    ["شیراز", "1404-07-03", 30],
    ["شیراز", "1404-08-01", 40],
    ["مشهد", "1404-08-02", 50],
    ["یزد", "1404-08-03", 300],
]


def test_report_numbers_match_an_independent_calculation() -> None:
    view = present(analyze(_table(FULL_ROWS), "مبلغ"), "report")
    amounts = [Decimal(row[2]) for row in FULL_ROWS[1:]]
    assert sum(amounts, Decimal(0)) == Decimal(450)
    assert _card(view, "total") == "450"
    assert _card(view, "mean") == "75"
    assert _card(view, "median") == "35"
    assert _card(view, "minimum") == "10"
    assert _card(view, "maximum") == "300"
    assert _card(view, "top_city") == "یزد"
    assert _card(view, "top_total") == "300"
    assert _card(view, "growth") == "550"
    assert _card(view, "outlier_count") == "1"
    assert next(card["hypothesis"] for card in view["cards"] if card["name"] == "outlier_count") is True
    assert consistent(view["summary"], view["metrics"])
    assert "99" not in view["summary"]
    periods = _sheet(view, "دوره‌ها")
    assert ["1404-07", "60"] in periods
    assert ["1404-08", "390"] in periods
    outliers = _sheet(view, "پرت")
    assert outliers == [["300", "احتمال"]]
    assert all(row[-1] != "خطا" for row in outliers)


def test_missing_column_does_not_invent_a_card() -> None:
    view = present(analyze(_table([["نام", "مبلغ"], ["علی", 10], ["سارا", 20]]), "مبلغ"), "report")
    names = {card["name"] for card in view["cards"]}
    assert "top_city" not in names
    assert "top_total" not in names
    assert "growth" not in names
    assert "outlier_count" not in names
    assert _card(view, "total") == "30"
    limits = _sheet(view, "محدودیت‌ها")
    assert ["رتبه", "قابل محاسبه نیست", "شهر"] in limits
    assert ["روند", "قابل محاسبه نیست", "تاریخ"] in limits
    assert ["پرت", "قابل محاسبه نیست", "نمونه"] in limits
    assert all(row[1] == "قابل محاسبه نیست" for row in limits)
    assert consistent(view["summary"], view["metrics"])


def test_a_number_outside_the_metrics_is_rejected() -> None:
    metrics = {"total": "450"}
    assert consistent("مجموع 450 است.", metrics)
    assert not consistent("مجموع 450 و حاشیه 99 است.", metrics)
    assert not consistent("حاشیه ۹۹ است.", metrics)
    with pytest.raises(ValueError):
        render("مجموع {total} و حاشیه 99.", metrics)


def test_rank_sentence_uses_the_city_total_and_keeps_the_original(client: TestClient) -> None:
    plan = build_plan(RANK_TEXT, "w", ["f"])
    assert plan["goal"] == "rank"
    assert all(step["tool"] != "chart" for step in plan["steps"])
    payload = _xlsx(FULL_ROWS)
    workspace, stored = _one(client, "رتبه", "فروش.xlsx", payload)
    response = client.post(
        f"/api/workspaces/{workspace['id']}/operations",
        json={"text": RANK_TEXT, "file_ids": [stored["id"]]},
    )
    assert response.status_code == 201, response.text
    body = response.json()
    assert body["status"] == "succeeded"
    assert body["plan"]["goal"] == "rank"
    assert "یزد" in body["summary"]
    assert "300" in body["summary"]
    assert body["error"] is None
    assert _card_from(body, "top_city") == "یزد"
    assert _card_from(body, "top_total") == "300"
    output = _download(client, workspace["id"], body["outputs"][0]["version_id"])
    book = openpyxl.load_workbook(BytesIO(output))
    assert book.sheetnames == ["شاخص‌ها", "شهرها"]
    assert book["شاخص‌ها"].sheet_view.rightToLeft is True
    cities = list(book["شهرها"].iter_rows(values_only=True))
    assert ["یزد", "300", "66.6667"] in [list(row) for row in cities]
    book.close()
    _assert_original(client, workspace["id"], stored, payload)


def test_report_sentence_omits_a_missing_city_card(client: TestClient) -> None:
    payload = _xlsx([["نام", "مبلغ"], ["علی", 10], ["سارا", 20]])
    workspace, stored = _one(client, "گزارش", "فروش.xlsx", payload)
    response = client.post(
        f"/api/workspaces/{workspace['id']}/operations",
        json={"text": REPORT_TEXT, "file_ids": [stored["id"]]},
    )
    assert response.status_code == 201, response.text
    body = response.json()
    assert body["summary"].startswith("مجموع 30")
    assert "شهر برتر" not in response.text
    assert all(card["name"] != "top_city" for card in body["cards"])
    output = _download(client, workspace["id"], body["outputs"][0]["version_id"])
    book = openpyxl.load_workbook(BytesIO(output))
    assert "محدودیت‌ها" in book.sheetnames
    assert "شهرها" not in book.sheetnames
    limits = list(book["محدودیت‌ها"].iter_rows(values_only=True))
    assert ("رتبه", "قابل محاسبه نیست", "شهر") in limits
    book.close()
    assert "گزارش.xlsx" in _names(client, workspace["id"])
    _assert_original(client, workspace["id"], stored, payload)


def test_rank_without_a_city_writes_nothing(client: TestClient) -> None:
    payload = _xlsx([["نام", "مبلغ"], ["علی", 60_000_000]])
    workspace, stored = _one(client, "بدون شهر", "فروش.xlsx", payload)
    before = _names(client, workspace["id"])
    response = client.post(
        f"/api/workspaces/{workspace['id']}/operations",
        json={"text": RANK_TEXT, "file_ids": [stored["id"]]},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["status"] == "waiting_confirmation"
    assert body["outputs"] == []
    assert body["cards"] == []
    assert "ستون شهر" in body["summary"]
    assert _names(client, workspace["id"]) == before


def _table(rows: list[list]) -> dict:
    headers = rows[0]
    roles = {"شهر": "city", "تاریخ": "date", "مبلغ": "amount", "نام": "person_name"}
    return {
        "sheet": "فروش",
        "headers": headers,
        "roles": {name: roles.get(name) for name in headers},
        "rows": [dict(zip(headers, row, strict=True)) for row in rows[1:]],
    }


def _card(view: dict, name: str) -> str:
    return next(card["value"] for card in view["cards"] if card["name"] == name)


def _card_from(body: dict, name: str) -> str:
    return next(card["value"] for card in body["cards"] if card["name"] == name)


def _sheet(view: dict, name: str) -> list[list[str]]:
    sheet = next(item for item in view["sheets"] if item["name"] == name)
    return [row[1:] if name == "پرت" else row for row in sheet["rows"][1:]]
