"""ChartSpec از جدول واقعی. سری ثابت ندارد."""

from __future__ import annotations

import re
from decimal import Decimal
from io import BytesIO

import openpyxl

from hesabdar.errors import AppError
from hesabdar.inspection.classify import (
    classify_date,
    compact_header,
    normalize_letters,
    parse_number,
    role_from_header,
    translate_digits,
)
from hesabdar.messages import OPERATION_INVALID, PLAN_NO_AMOUNT, PLAN_NO_CATEGORY, PLAN_NO_VALUES

CHART_TYPES = ("column", "bar", "line", "pie", "scatter", "area", "combo")


def grouping_columns(table: dict, role: str | None = None) -> list[str]:
    found = []
    for name in table["headers"]:
        if _matches(table, name, role):
            found.append(name)
    return found


def build_chart(
    table: dict,
    amount_column: str,
    chart_type: str,
    title: str,
    category_column: str | None,
    source: str,
) -> dict:
    if chart_type not in CHART_TYPES:
        return _blocked(PLAN_NO_CATEGORY)
    if chart_type == "scatter":
        points, seen, invalid = _scatter(table, amount_column)
        category_field = "ردیف"
    else:
        if not category_column:
            return _blocked(PLAN_NO_CATEGORY)
        points, seen, invalid = _aggregate(table, amount_column, category_column)
        category_field = category_column
    if not seen:
        return _blocked(PLAN_NO_VALUES if invalid else PLAN_NO_AMOUNT)
    if not points:
        return _blocked(PLAN_NO_CATEGORY)
    spec = make_spec(chart_type, title, category_field, points, amount_column, source)
    return {"blocked": None, "spec": spec, "sheet": sheet_for(spec)}


def make_spec(
    chart_type: str,
    title: str,
    category_field: str,
    points: list[dict],
    series_name: str,
    source: str,
) -> dict:
    series = [{"name": series_name, "value_field": series_name, "points": points}]
    if chart_type == "combo":
        overlay = "روند خطی" if series_name == "روند" else "روند"
        series.append({"name": overlay, "value_field": series_name, "points": points})
    return {
        "chart_type": chart_type,
        "title": title,
        "category_field": category_field,
        "series": series,
        "source": source,
    }


def sheet_for(spec: dict, sheet_name: str = "نمودار") -> dict:
    headers = [spec["category_field"]]
    headers.extend(series["name"] for series in spec["series"])
    rows = [headers]
    numbers: set[tuple[int, int]] = set()
    for index, point in enumerate(spec["series"][0]["points"]):
        row_index = index + 1
        row = [point["category"]]
        if spec["chart_type"] == "scatter":
            numbers.add((row_index, 0))
        for series_index, series in enumerate(spec["series"]):
            row.append(series["points"][index]["value"])
            numbers.add((row_index, series_index + 1))
        rows.append(row)
    return {
        "name": sheet_name,
        "rows": rows,
        "numbers": numbers,
        "chart": {
            "type": spec["chart_type"],
            "title": spec["title"],
            "anchor": "E2",
            "max_row": len(rows),
            "category_col": 1,
            "series_cols": list(range(2, len(headers) + 1)),
            "series_title": spec["series"][0]["name"],
        },
    }


def assert_charts(data: bytes, pairs: list[tuple[str, dict]]) -> None:
    try:
        book = openpyxl.load_workbook(BytesIO(data))
    except Exception:
        raise AppError("output.invalid", OPERATION_INVALID, 422) from None
    try:
        for sheet_name, spec in pairs:
            if sheet_name not in book.sheetnames:
                raise AppError("output.invalid", OPERATION_INVALID, 422)
            sheet = book[sheet_name]
            charts = getattr(sheet, "_charts", [])
            if not charts or read_chart_title(charts[0]) != spec["title"]:
                raise AppError("output.invalid", OPERATION_INVALID, 422)
            body = list(sheet.iter_rows(min_row=2, values_only=True))
            points = spec["series"][0]["points"]
            if len(body) != len(points):
                raise AppError("output.invalid", OPERATION_INVALID, 422)
            for index, point in enumerate(points):
                if str(body[index][0]) != str(point["category"]):
                    raise AppError("output.invalid", OPERATION_INVALID, 422)
                for series_index, series in enumerate(spec["series"]):
                    if not _same(body[index][series_index + 1], series["points"][index]["value"]):
                        raise AppError("output.invalid", OPERATION_INVALID, 422)
    finally:
        book.close()


def read_chart_title(chart) -> str:
    title = getattr(chart, "title", None)
    if title is None:
        return ""
    if isinstance(title, str):
        return title
    text = getattr(title, "tx", None)
    if isinstance(text, str):
        return text
    rich = getattr(text, "rich", None) if text is not None else None
    if rich is None:
        return ""
    parts = []
    for paragraph in rich.p or []:
        for run in paragraph.r or []:
            if run.t:
                parts.append(str(run.t))
        if not paragraph.r and getattr(paragraph, "endParaRPr", None) is None:
            continue
    return "".join(parts)


def plain(value: Decimal) -> str:
    rendered = format(value, "f")
    if "." in rendered:
        rendered = rendered.rstrip("0").rstrip(".")
    return rendered


def _blocked(message: str) -> dict:
    return {"blocked": message, "spec": None, "sheet": None}


def _matches(table: dict, name: str, role: str | None) -> bool:
    kind = table["roles"].get(name) or role_from_header(name)
    if role == "city":
        return kind == "city"
    if role == "date":
        return kind == "date"
    if role == "product":
        return _is_product(name)
    if role == "category":
        return _is_category(name)
    return kind in {"city", "date"} or _is_product(name) or _is_category(name)


def _is_product(name: str) -> bool:
    compact = compact_header(name)
    if "کد" in compact:
        return False
    return "محصول" in compact or "کالا" in compact


def _is_category(name: str) -> bool:
    compact = compact_header(name)
    return "دسته" in compact or "گروه" in compact


def _aggregate(table: dict, amount_column: str, category_column: str):
    kind = table["roles"].get(category_column) or role_from_header(category_column)
    as_period = kind == "date"
    order: list[str] = []
    totals: dict[str, Decimal] = {}
    labels: dict[str, str] = {}
    seen = False
    invalid = 0
    for row in table["rows"]:
        amount, bad = _amount(row.get(amount_column))
        if bad:
            invalid += 1
            continue
        if amount is None:
            continue
        seen = True
        parsed = _period(row.get(category_column)) if as_period else _label(row.get(category_column))
        if parsed is None:
            continue
        key, label = parsed
        if key not in totals:
            order.append(key)
            totals[key] = Decimal(0)
            labels[key] = label
        totals[key] += amount
    points = [{"category": labels[key], "value": plain(totals[key])} for key in order]
    return points, seen, invalid


def _scatter(table: dict, amount_column: str):
    points = []
    seen = False
    invalid = 0
    for index, row in enumerate(table["rows"], start=1):
        amount, bad = _amount(row.get(amount_column))
        if bad:
            invalid += 1
            continue
        if amount is None:
            continue
        seen = True
        points.append({"category": str(index), "value": plain(amount)})
    return points, seen, invalid


def _amount(raw) -> tuple[Decimal | None, bool]:
    if raw in (None, ""):
        return None, False
    number, warning = parse_number(str(raw))
    if warning or number is None:
        return None, True
    return number, False


def _label(raw):
    if raw in (None, ""):
        return None
    label = str(raw).strip()
    if not label:
        return None
    key = re.sub(r"\s+", " ", translate_digits(normalize_letters(label))).strip()
    return key, label


def _period(raw):
    if raw in (None, ""):
        return None
    parsed, warning = classify_date(str(raw))
    if warning or not parsed or len(parsed) < 7:
        return None
    period = parsed[:7]
    return period, period


def _same(cell, expected: str) -> bool:
    if cell is None:
        return False
    try:
        return Decimal(str(cell)) == Decimal(expected)
    except Exception:
        return str(cell) == expected
