from __future__ import annotations

import csv
from io import BytesIO, StringIO

from decimal import Decimal

import xlwt
from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.chart import AreaChart, BarChart, LineChart, PieChart, Reference, ScatterChart, Series
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from hesabdar.errors import AppError
from hesabdar.messages import FILE_UNSUPPORTED


def write_workbook(sheets: list[dict], file_format: str) -> bytes:
    if file_format == "xlsx":
        return _write_xlsx(sheets)
    if file_format == "xls":
        return _write_xls(sheets)
    if file_format == "csv":
        return _write_csv(sheets)
    raise AppError("file.unsupported_type", FILE_UNSUPPORTED, 400)


def _write_xlsx(sheets: list[dict]) -> bytes:
    book = Workbook()
    first = book.active
    for index, sheet in enumerate(sheets):
        worksheet = first if index == 0 else book.create_sheet()
        worksheet.title = sheet["name"]
        worksheet.sheet_view.rightToLeft = True
        formulas = sheet.get("formulas") or set()
        numbers = sheet.get("numbers") or set()
        for row_index, row in enumerate(sheet["rows"]):
            for column_index, value in enumerate(row):
                if value is None:
                    continue
                cell = worksheet.cell(row_index + 1, column_index + 1)
                if (row_index, column_index) in numbers:
                    _write_number(cell, value)
                else:
                    _write_xlsx_cell(cell, value, (row_index, column_index) in formulas)
        if sheet.get("chart"):
            _add_chart(worksheet, sheet["chart"])
        _finish_sheet(worksheet, sheet)
    if not sheets:
        first.title = "خالی"
        first.sheet_view.rightToLeft = True
    buffer = BytesIO()
    book.save(buffer)
    book.close()
    return buffer.getvalue()


def _write_number(cell: Cell, value: str) -> None:
    number = Decimal(str(value))
    cell.value = int(number) if number == number.to_integral_value() else float(number)
    cell.number_format = "0.####"
    cell.font = Font(name="Tahoma")


def _add_chart(worksheet, spec: dict) -> None:
    max_row = spec["max_row"]
    categories = Reference(worksheet, min_col=spec["category_col"], min_row=2, max_row=max_row)
    chart_type = spec["type"]
    if chart_type == "scatter":
        chart = ScatterChart()
        x_values = Reference(worksheet, min_col=spec["category_col"], min_row=2, max_row=max_row)
        y_values = Reference(worksheet, min_col=spec["series_cols"][0], min_row=2, max_row=max_row)
        series = Series(y_values, x_values, title=spec.get("series_title") or "مبلغ")
        chart.series.append(series)
    elif chart_type == "combo":
        chart = BarChart()
        chart.type = "col"
        data = Reference(worksheet, min_col=spec["series_cols"][0], min_row=1, max_row=max_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)
        line = LineChart()
        line_data = Reference(worksheet, min_col=spec["series_cols"][1], min_row=1, max_row=max_row)
        line.add_data(line_data, titles_from_data=True)
        line.set_categories(categories)
        chart.y_axis.crosses = "min"
        line.y_axis.axId = 200
        chart += line
    else:
        chart = _chart_object(chart_type)
        data = Reference(worksheet, min_col=spec["series_cols"][0], min_row=1, max_row=max_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)
    chart.title = spec["title"]
    worksheet.add_chart(chart, spec.get("anchor") or "E2")


def _chart_object(chart_type: str):
    if chart_type == "column":
        chart = BarChart()
        chart.type = "col"
        return chart
    if chart_type == "bar":
        chart = BarChart()
        chart.type = "bar"
        return chart
    if chart_type == "line":
        return LineChart()
    if chart_type == "pie":
        return PieChart()
    if chart_type == "area":
        return AreaChart()
    raise AppError("file.unsupported_type", FILE_UNSUPPORTED, 400)


def _write_xlsx_cell(cell: Cell, value: str, formula: bool) -> None:
    cell.value = value
    cell.font = Font(name="Tahoma")
    if formula:
        return
    cell.number_format = "@"
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"


def _finish_sheet(worksheet, sheet: dict) -> None:
    if sheet.get("bold_header"):
        for cell in worksheet[1]:
            if cell.value is not None:
                cell.font = Font(name="Tahoma", bold=True)
    if sheet.get("freeze"):
        worksheet.freeze_panes = sheet["freeze"]
    outline = sheet.get("outline")
    if outline and int(outline["end"]) >= int(outline["start"]):
        worksheet.row_dimensions.group(int(outline["start"]), int(outline["end"]), hidden=False)
    for column in sheet.get("text_columns") or []:
        worksheet.column_dimensions[get_column_letter(int(column))].number_format = "@"
    headers = sheet["rows"][0] if sheet.get("rows") else []
    for index, header in enumerate(headers, start=1):
        letter = get_column_letter(index)
        worksheet.column_dimensions[letter].width = max(14, min(36, len(str(header)) + 6))
    validation = sheet.get("validation")
    if not validation:
        return
    formula = '"' + ",".join(str(value) for value in validation["values"]) + '"'
    rule = DataValidation(type="list", formula1=formula, allow_blank=True)
    rule.add(validation["sqref"])
    worksheet.add_data_validation(rule)


def _write_xls(sheets: list[dict]) -> bytes:
    book = xlwt.Workbook(encoding="utf-8")
    if not sheets:
        book.add_sheet("خالی")
    for sheet in sheets:
        worksheet = book.add_sheet(sheet["name"][:31] or "برگه")
        for row_index, row in enumerate(sheet["rows"]):
            for column_index, value in enumerate(row):
                if value is not None:
                    worksheet.write(row_index, column_index, value)
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _write_csv(sheets: list[dict]) -> bytes:
    rows = sheets[0]["rows"] if sheets else []
    buffer = StringIO(newline="")
    writer = csv.writer(buffer, lineterminator="\n")
    for row in rows:
        writer.writerow(["" if value is None else value for value in row])
    return buffer.getvalue().encode("utf-8-sig")
