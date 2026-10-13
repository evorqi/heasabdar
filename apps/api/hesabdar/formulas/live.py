"""فرمول زنده با ارجاع صریح. مقدار کش‌شده اکسل ساخته نمی‌شود."""

from __future__ import annotations

from io import BytesIO

import openpyxl

from hesabdar.errors import AppError
from hesabdar.messages import OPERATION_INVALID


def column_letter(index: int) -> str:
    number = index + 1
    letters = ""
    while number:
        number, remainder = divmod(number - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def percent_formula(base_index: int, other_index: int, excel_row: int) -> str:
    base = column_letter(base_index)
    other = column_letter(other_index)
    return f'=IF({base}{excel_row}=0,"",ABS({other}{excel_row}-{base}{excel_row})/ABS({base}{excel_row})*100)'


def assert_live_formulas(data: bytes, expected: list[tuple[int, int, str]]) -> None:
    try:
        book = openpyxl.load_workbook(BytesIO(data))
        cached = openpyxl.load_workbook(BytesIO(data), data_only=True)
    except Exception:
        raise AppError("output.invalid", OPERATION_INVALID, 422) from None
    try:
        sheet = book.worksheets[0]
        cached_sheet = cached.worksheets[0]
        for row, column, text in expected:
            if sheet.cell(row, column).value != text:
                raise AppError("output.invalid", OPERATION_INVALID, 422)
            if cached_sheet.cell(row, column).value is not None:
                raise AppError("output.invalid", OPERATION_INVALID, 422)
        if "درباره این خروجی" not in book.sheetnames:
            raise AppError("output.invalid", OPERATION_INVALID, 422)
        about = list(book["درباره این خروجی"].iter_rows(values_only=True))
        if ("فرمول زنده", "ارزیابی نشده") not in about:
            raise AppError("output.invalid", OPERATION_INVALID, 422)
    finally:
        book.close()
        cached.close()
